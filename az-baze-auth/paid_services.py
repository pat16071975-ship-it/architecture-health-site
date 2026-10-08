import re
from datetime import date, datetime
from io import BytesIO

from openpyxl import load_workbook


STAFF_HEADER_RE = re.compile(
    r"^(?:[A-Za-zА-ЯЁа-яё-]+(?:\s+[A-Za-zА-ЯЁа-яё-]+)*)\s+[A-ZА-ЯЁ]\.\s*(?:[A-ZА-ЯЁ]\.)?$"
)
INVOICE_RE = re.compile(r"^Счет №(\d+) от (\d{2}\.\d{2}\.\d{4})(?:\s+\d{1,2}:\d{2}:\d{2})?$")
DATE_TIME_FORMATS = ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%d.%m.%Y")


def _text(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _num(value):
    if value is None:
        return 0.0
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace("\xa0", " ").replace("₽", "").strip()
    if text in {"", "-", "—"}:
        return 0.0
    text = text.replace(" ", "").replace(",", ".")
    try:
        return float(text)
    except ValueError as exc:
        raise ValueError(f"Не удалось прочитать денежное значение «{value}».") from exc


def _qty(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = str(value).strip().replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def _iso(value):
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    text = _text(value)
    for fmt in DATE_TIME_FORMATS:
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return None


def _money_row(values):
    return {
        "opening": round(_num(values[0]), 2),
        "billed": round(_num(values[1]), 2),
        "paid": round(_num(values[2]), 2),
        "closing": round(_num(values[3]), 2),
    }


def _assert_balance(label, values, tolerance=0.02):
    delta = round(values["opening"] + values["billed"] - values["paid"] - values["closing"], 2)
    if abs(delta) > tolerance:
        raise ValueError(
            f"Нарушен баланс нового отчёта для «{label}»: "
            "задолженность на начало + сумма со скидкой - оплачено "
            "не равны задолженности на конец."
        )
    return delta


def _find_sheet(workbook):
    required = {
        "Группа услуг",
        "Услуги",
        "Задолженность на нач. периода",
        "Сумма со скидкой",
        "Оплачено",
        "Задолженность на конец периода",
    }
    for ws in workbook.worksheets:
        headers = {_text(cell.value) for cell in ws[1] if _text(cell.value)}
        if required.issubset(headers):
            return ws
    raise ValueError(
        "Не найден лист нового отчёта «Выручка по направлениям» "
        "с колонками задолженности и оплаты."
    )


def _provider_header(row):
    first = _text(row[0])
    return bool(
        first
        and STAFF_HEADER_RE.fullmatch(first)
        and not _text(row[1])
        and not _text(row[2])
        and not _text(row[3])
    )


def parse_bytes(raw, filename=""):
    if not raw:
        raise ValueError("Выбран пустой файл «Выручка по направлениям».")
    if len(raw) > 25 * 1024 * 1024:
        raise ValueError("Размер файла «Выручка по направлениям» превышает 25 МБ.")

    try:
        wb = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("Не удалось открыть новый отчёт «Выручка по направлениям».") from exc

    try:
        ws = _find_sheet(wb)
        section_found = False
        section_totals = None
        providers = {}
        items = []
        invoices = []
        current_staff = None
        current_patient = ""
        current_invoice = None
        current_date = None
        current_group = ""

        for row in ws.iter_rows(values_only=True):
            values = list(row[:8]) + [None] * max(0, 8 - len(row[:8]))
            first = _text(values[0])
            second = _text(values[1])
            service = _text(values[2])
            qty = _qty(values[3])

            if not section_found:
                if first == "Услуги":
                    section_found = True
                    section_totals = _money_row(values[4:8])
                    _assert_balance("Услуги", section_totals)
                continue

            if first == "Авансы":
                break

            if first == "Неоплаченные услуги на начало периода":
                # These rows explain opening debt. Their payment is already
                # included in the provider-level «Оплачено» total and must not
                # become a second service item.
                continue

            if _provider_header(values):
                current_staff = first
                current_patient = ""
                current_invoice = None
                current_date = None
                current_group = ""
                summary = _money_row(values[4:8])
                _assert_balance(current_staff, summary)
                providers[current_staff] = summary
                continue

            if not current_staff:
                continue

            match = INVOICE_RE.match(first)
            if match:
                current_invoice = match.group(1)
                current_date = datetime.strptime(match.group(2), "%d.%m.%Y").strftime("%Y-%m-%d")
                current_group = ""
                invoices.append(
                    {
                        "staff": current_staff,
                        "invoice": current_invoice,
                        "date": current_date,
                    }
                )
                continue

            # Patient header between staff and invoice rows.
            if first and not second and not service and qty is None:
                current_patient = first
                continue

            if not service or qty is None:
                continue

            service_date = _iso(values[1]) or current_date
            if not service_date:
                continue

            if first:
                current_group = first

            billed = _num(values[5])
            paid = _num(values[6])
            closing = _num(values[7])
            items.append(
                {
                    "staff": current_staff,
                    "patient": current_patient,
                    "date": service_date,
                    "group": current_group,
                    "service": service,
                    "qty": qty,
                    "amount": round(billed, 2),
                    "paid_amount": round(paid, 2),
                    "closing_debt": round(closing, 2),
                    "invoice": current_invoice,
                }
            )

        if not section_found or section_totals is None:
            raise ValueError("В новом отчёте не найден раздел «Услуги».")
        if not providers:
            raise ValueError("В новом отчёте не найдены строки врачей/исполнителей.")
        if not items:
            raise ValueError("В новом отчёте не найдены строки услуг выбранного периода.")

        months = sorted({str(item["date"])[:7] for item in items if item.get("date")})
        if len(months) != 1:
            raise ValueError("Новый отчёт «Выручка по направлениям» должен содержать один календарный месяц.")
        month = months[0]
        period_end = max(str(item["date"]) for item in items if item.get("date"))

        provider_sum = {
            key: round(sum(value[key] for value in providers.values()), 2)
            for key in ("opening", "billed", "paid", "closing")
        }
        residual = {
            key: round(section_totals[key] - provider_sum[key], 2)
            for key in ("opening", "billed", "paid", "closing")
        }
        _assert_balance("нераспознанная часть отчёта", residual)

        return {
            "source_type": "paid_services_v1",
            "sheet": ws.title,
            "filename": filename or "",
            "month": month,
            "period_end": period_end,
            "totals": section_totals,
            "providers": providers,
            "provider_residual": residual,
            "items": items,
            "invoices": invoices,
        }
    finally:
        wb.close()


def direction_summary(report, dentists, structure_doctors, lab_doctors, ignored=None):
    ignored = set(ignored or ())
    providers = report.get("providers") or {}
    dentists = dict(dentists or {})
    structure_doctors = dict(structure_doctors or {})
    lab_doctors = dict(lab_doctors or {})

    dent_rows = {}
    structure_rows = {}
    lab_rows = {}
    unknown = {}
    ignored_paid = 0.0

    for source_name, values in providers.items():
        paid = round(_num((values or {}).get("paid")), 2)
        if source_name in dentists:
            display = dentists[source_name]
            dent_rows[display] = round(dent_rows.get(display, 0.0) + paid, 2)
        elif source_name in structure_doctors:
            display = structure_doctors[source_name]
            structure_rows[display] = round(structure_rows.get(display, 0.0) + paid, 2)
        elif source_name in lab_doctors:
            display = lab_doctors[source_name]
            lab_rows[display] = round(lab_rows.get(display, 0.0) + paid, 2)
        elif source_name in ignored:
            ignored_paid = round(ignored_paid + paid, 2)
        elif abs(paid) > 0.004 or abs(_num((values or {}).get("billed"))) > 0.004:
            unknown[source_name] = {
                "paid": paid,
                "billed": round(_num((values or {}).get("billed")), 2),
                "opening": round(_num((values or {}).get("opening")), 2),
                "closing": round(_num((values or {}).get("closing")), 2),
            }

    dent_total = round(sum(dent_rows.values()), 2)
    structure_total = round(sum(structure_rows.values()), 2)
    lab_total = round(sum(lab_rows.values()), 2)
    residual_paid = round(_num((report.get("provider_residual") or {}).get("paid")), 2)
    paid_total = round(_num((report.get("totals") or {}).get("paid")), 2)
    classified_total = round(dent_total + structure_total + lab_total, 2)

    return {
        "dentists": dent_rows,
        "clinicDocs": structure_rows,
        "labDocs": lab_rows,
        "dentPaid": dent_total,
        "structurePaid": structure_total,
        "labPaid": lab_total,
        "classifiedPaid": classified_total,
        "paidTotal": paid_total,
        "unknownProviders": unknown,
        "unknownPaid": round(sum(item["paid"] for item in unknown.values()), 2),
        "ignoredPaid": ignored_paid,
        "nonProviderPaid": residual_paid,
        "unclassifiedPaid": round(paid_total - classified_total, 2),
    }
