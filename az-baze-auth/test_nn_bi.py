import io
import json
import os
import sqlite3
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

os.environ.setdefault("AZBAZE_SECRET_KEY", "test-secret")
os.environ["AZBAZE_OWNER_EMAIL"] = "owner@example.com"

import app as app_module
import nn_normalize
import nn_upload


class NNBIAnalyticsTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        root = Path(self.tmp.name)
        app_module.DB_PATH = root / "auth.db"
        app_module.SITE_ROOT = root / "site"
        nn_upload.DB_PATH = app_module.DB_PATH
        nn_upload.UPLOAD_ROOT = root / "nn-uploads"
        app_module.SITE_ROOT.mkdir(parents=True, exist_ok=True)
        (app_module.SITE_ROOT / "index.html").write_text(
            "<!doctype html><html><body><nav></nav></body></html>", encoding="utf-8"
        )
        app_module.init_db_file(app_module.DB_PATH)

        conn = sqlite3.connect(app_module.DB_PATH)
        conn.row_factory = sqlite3.Row
        conn.execute("PRAGMA foreign_keys=ON")
        conn.executescript(
            """
            CREATE TABLE clinics (
                id INTEGER PRIMARY KEY,
                name TEXT NOT NULL,
                address TEXT NOT NULL DEFAULT '',
                status TEXT NOT NULL
            );
            """
        )
        now = app_module.iso_now()
        conn.executemany(
            """INSERT INTO users(
                id,email,full_name,password_hash,is_admin,active,must_change_password,
                failed_attempts,locked_until,created_at,updated_at
            ) VALUES(?,?,?,?,?,1,0,0,NULL,?,?)""",
            [
                (1, "owner@example.com", "Owner", "x", 1, now, now),
                (2, "reports@example.com", "Reports", "x", 0, now, now),
                (3, "upload@example.com", "Upload", "x", 0, now, now),
            ],
        )
        conn.executemany(
            "INSERT INTO permissions(user_id,section) VALUES(?,?)",
            [(2, "nn_reports"), (3, "nn_upload")],
        )
        conn.execute(
            "INSERT INTO clinics(id,name,address,status) VALUES(11,'Архитектура здоровья','AZ address','active')"
        )
        conn.commit()
        conn.close()

        nn_upload._init_schema()
        conn = sqlite3.connect(app_module.DB_PATH)
        conn.row_factory = sqlite3.Row
        conn.execute(
            "INSERT INTO nn_upload_batches(id,clinic_id,bundle_sha256,status,uploaded_by,uploaded_at) VALUES(101,1,'hash','ready',1,?)",
            (now,),
        )
        conn.commit()
        nn_normalize._init_schema(conn)

        payload = {
            "version": 1,
            "period": {"start": "2026-06-01", "end": "2026-06-30"},
            "source_control": {},
            "patients": [
                {"key": "card:1", "name": "Пациент Один", "dob": "1980-01-01", "gender": "Женский", "source": "Instagram", "visits_count": 3, "iin": "", "chart": "1", "phone": "7000000001", "note": ""},
                {"key": "card:2", "name": "Пациент Два", "dob": "1995-06-15", "gender": "Мужской", "source": "2GIS", "visits_count": 1, "iin": "", "chart": "2", "phone": "7000000002", "note": ""},
                {"key": "card:3", "name": "Пациент Три", "dob": "1960-03-20", "gender": "Женский", "source": "Instagram", "visits_count": 2, "iin": "", "chart": "3", "phone": "7000000003", "note": ""},
            ],
            "visits": [
                {"patient_key": "card:1", "date": "2026-06-01", "created_at": "2026-05-25T10:00", "start": "09:00", "end": "09:30", "doctor": "Врач А", "patient": "Пациент Один", "chart": "1", "phone": "7000000001", "iin": "", "repeat_field": "", "note": "ЭВЛК рекомендовано", "visit_type": "Амбулаторно", "help_type": "", "appeal_reason": "", "diagnoses": "I83.9", "services_text": "Прием хирурга флеболога с УЗИ", "amount": 6000},
                {"patient_key": "card:1", "date": "2026-06-05", "created_at": "2026-06-01T11:00", "start": "08:00", "end": "10:00", "doctor": "Врач А", "patient": "Пациент Один", "chart": "1", "phone": "7000000001", "iin": "", "repeat_field": "", "note": "ЭВЛК выполнено", "visit_type": "Амбулаторно", "help_type": "", "appeal_reason": "", "diagnoses": "I83.9", "services_text": "Лазерное лечение варикоза категория B", "amount": 350000},
                {"patient_key": "card:1", "date": "2026-06-12", "created_at": "2026-06-05T12:00", "start": "10:00", "end": "10:30", "doctor": "Врач А", "patient": "Пациент Один", "chart": "1", "phone": "7000000001", "iin": "", "repeat_field": "повторный", "note": "контроль после операции", "visit_type": "Амбулаторно", "help_type": "", "appeal_reason": "", "diagnoses": "", "services_text": "", "amount": 0},
                {"patient_key": "card:2", "date": "2026-06-02", "created_at": "2026-05-30T09:00", "start": "14:00", "end": "14:30", "doctor": "Врач Б", "patient": "Пациент Два", "chart": "2", "phone": "7000000002", "iin": "", "repeat_field": "", "note": "", "visit_type": "Амбулаторно", "help_type": "", "appeal_reason": "", "diagnoses": "", "services_text": "Прием хирурга флеболога с УЗИ", "amount": 0},
                {"patient_key": "card:3", "date": "2026-06-03", "created_at": "2026-05-20T09:00", "start": "17:30", "end": "18:00", "doctor": "Врач Б", "patient": "Пациент Три", "chart": "3", "phone": "7000000003", "iin": "", "repeat_field": "", "note": "склеротерапия рекомендована", "visit_type": "Амбулаторно", "help_type": "", "appeal_reason": "", "diagnoses": "", "services_text": "Прием хирурга флеболога с УЗИ", "amount": 6000},
                {"patient_key": "card:3", "date": "2026-06-20", "created_at": "2026-06-10T09:00", "start": "18:00", "end": "19:00", "doctor": "Врач Б", "patient": "Пациент Три", "chart": "3", "phone": "7000000003", "iin": "", "repeat_field": "", "note": "пенная склеротерапия", "visit_type": "Амбулаторно", "help_type": "", "appeal_reason": "", "diagnoses": "", "services_text": "Пенная склеротерапия варикозных вен", "amount": 150000},
            ],
            "medical_records": [
                {"date": "2026-06-05", "doctor": "Врач А", "patient": "Пациент Один", "dob": "1980-01-01", "chart": "1", "complaint": "", "objective": "", "direction": "", "assignment": "", "recommendation": "", "treatment": "Протокол манипуляции: Эндоваскулярная лазерная коагуляция"},
                {"date": "2026-06-20", "doctor": "Врач Б", "patient": "Пациент Три", "dob": "1960-03-20", "chart": "3", "complaint": "", "objective": "", "direction": "", "assignment": "", "recommendation": "", "treatment": "Протокол манипуляции: Пенная Склеротерапия"},
            ],
            "service_control_rows": [
                {"service": "Лазерное лечение варикоза категория B", "qty": 1, "gross_amount": 367000, "discount": 17000, "amount": 350000, "aggregate_control": True},
                {"service": "Пенная склеротерапия", "qty": 1, "gross_amount": 150000, "discount": 0, "amount": 150000, "aggregate_control": True},
            ],
            "payment_control_rows": [
                {"doctor": "Врач А", "method": "Безналичные", "qty": 2, "gross_amount": 373000, "discount": 17000, "amount": 356000, "period_start": "2026-06-01", "period_end": "2026-06-20"},
                {"doctor": "Врач Б", "method": "Наличные", "qty": 2, "gross_amount": 156000, "discount": 0, "amount": 156000, "period_start": "2026-06-01", "period_end": "2026-06-20"},
            ],
            "deleted_appointments": [
                {"deleted_by": "Администратор", "reason": "Неявка на прием", "deleted_date": "2026-06-10", "deleted_time": "12:00", "appointment_date": "2026-06-11", "appointment_time": "10:00", "doctor": "Врач А", "patient": "Пациент X", "patient_name": "пациент x", "phone": ""}
            ],
            "primaries": [], "repeats": [], "treatments": [], "suspicious": [], "monthly": [], "doctor_metrics": [], "key_metrics": {},
        }
        conn.execute(
            "INSERT INTO nn_normalized_batches(batch_id,clinic_id,version,period_start,period_end,payload_json,normalized_at) VALUES(101,1,1,'2026-06-01','2026-06-30',?,?)",
            (json.dumps(payload, ensure_ascii=False), now),
        )
        conn.commit()
        conn.close()

        self.web = app_module.create_app()
        self.web.config.update(TESTING=True)

    def tearDown(self):
        self.tmp.cleanup()

    def client_for(self, user_id):
        client = self.web.test_client()
        with client.session_transaction() as session:
            session["user_id"] = user_id
            session["csrf"] = "test-csrf"
        return client

    def test_home_exposes_bi_button_and_page_has_all_sections(self):
        home = self.client_for(2).get("/nn/")
        html = home.get_data(as_text=True)
        self.assertEqual(home.status_code, 200)
        self.assertIn("BI-аналитика", html)
        self.assertIn('/nn/bi/', html)
        self.assertIn("Графические отчеты", html)
        self.assertIn('/nn/graphics/', html)

        page = self.client_for(2).get("/nn/bi/")
        body = page.get_data(as_text=True)
        self.assertEqual(page.status_code, 200)
        for label in (
            "Обзор бизнеса", "Воронка", "Врачи", "Пациенты", "Услуги и лечение",
            "Запись и потери", "Выручка и скидки", "Маркетинг", "Связи",
        ):
            self.assertIn(label, body)
        self.assertIn('<a class="nn-back" href="/nn/">← Назад</a>', body)

    def test_empty_clinic_clears_all_bi_sections(self):
        response = self.client_for(2).get("/api/nn/bi/summary?clinic_id=2")
        self.assertEqual(response.status_code, 200)
        data = response.get_json()
        self.assertTrue(data["empty"])
        self.assertEqual(data["clinic_name"], "Клиника 2 (Шевчеко)")

        page = self.client_for(2).get("/nn/bi/")
        body = page.get_data(as_text=True)
        self.assertIn("function renderEmpty(d)", body)
        self.assertIn('["overview","funnel","doctors","patients","services","bookings","revenue","marketing","relations"]', body)
        self.assertIn('if(d.empty){renderEmpty(d);return}', body)
        self.assertIn('exportLink.removeAttribute("href")', body)
        self.assertIn('document.getElementById("relApply").disabled=true', body)

    def test_graphical_reports_page_uses_existing_bi_data_contract(self):
        page = self.client_for(2).get("/nn/graphics/")
        body = page.get_data(as_text=True)
        self.assertEqual(page.status_code, 200)
        self.assertIn("Графические отчеты", body)
        self.assertIn('/api/nn/bi/summary?', body)
        for label in (
            "Выручка во времени",
            "Пациенты во времени",
            "Приёмы во времени",
            "Выручка по врачам",
            "Конверсия первичный → лечение по врачам",
            "Выручка по категориям услуг",
            "Возраст пациентов",
            "Частота посещений",
            "Приёмы по дням недели",
            "Приёмы по времени суток",
            "Причины удалённых записей",
            "Скидки",
            "Способы оплаты",
            "Источники пациентов",
        ):
            self.assertIn(label, body)
        self.assertIn('data-size="third"', body)
        self.assertIn('data-size="half"', body)
        self.assertIn('data-size="full"', body)
        self.assertIn("function renderEmpty(d)", body)
        self.assertIn("font-size:22px;font-weight:700", body)
        self.assertIn("font-size:15px;line-height:1.45;font-weight:700", body)
        self.assertIn("function shortDoctor(value)", body)
        self.assertIn("function xLabelIndices(rows)", body)
        self.assertIn("rows.map((_,i)=>i)", body)
        self.assertIn('svgLine("revenue",d.overview.trend,"revenue",fmtMoney,true,true)', body)
        self.assertIn("H=compact?250:420", body)
        self.assertIn("gr-svg-compact", body)
        self.assertIn("{labelFormatter:shortDoctor}", body)
        self.assertIn('{labelsAbove:true}', body)
        self.assertIn('data-fill-toggle', body)
        self.assertIn("Заливка: Вкл", body)
        self.assertIn("grTooltip", body)
        self.assertIn("data-tip", body)
        self.assertIn('.gr-card[data-card="deleted"],.gr-card[data-card="discounts"]{height:285px', body)
        self.assertIn('.gr-card[data-card="deleted"] .gr-note,.gr-card[data-card="discounts"] .gr-note{font-size:13px', body)
        self.assertIn(".gr-svg-fit{min-width:0;width:100%;max-width:100%}", body)
        self.assertIn('svgBars("weekday",d.bookings.weekday,"label","visits",fmtNum,{compact:true,fit:true,smallText:true})', body)
        self.assertIn('svgBars("time",d.bookings.time,"label","visits",fmtNum,{compact:true,fit:true,smallText:true})', body)
        self.assertIn('svgBars("deleted",d.bookings.deleted_reasons,"label","value",fmtNum,{labelsAbove:true,smallText:true,tight:true})', body)
        self.assertIn('{compact:true,smallText:true,tight:true}', body)
        self.assertIn('<a class="nn-back" href="/nn/">← Назад</a>', body)

    def test_summary_covers_business_sides_and_filters(self):
        response = self.client_for(2).get(
            "/api/nn/bi/summary?clinic_id=1&date_from=2026-06-01&date_to=2026-06-30"
        )
        self.assertEqual(response.status_code, 200)
        data = response.get_json()
        self.assertFalse(data["empty"])
        # Requested end is clamped to the latest date actually available in the data.
        self.assertEqual(data["period"], {"start": "2026-06-01", "end": "2026-06-20"})
        self.assertEqual(data["overview"]["metrics"]["patients"], 3)
        self.assertGreater(data["overview"]["metrics"]["revenue"], 500000)
        self.assertGreaterEqual(data["funnel"]["conversion_primary_to_treatment"], 0)
        self.assertTrue(data["overview"]["trend"])
        self.assertTrue(data["doctors"])
        self.assertIn("doctor_transitions", data)
        self.assertIn("days_to_treatment", data["funnel"])
        self.assertTrue(data["patients"]["age_groups"])
        self.assertTrue(data["services"]["categories"])
        self.assertTrue(data["services"]["chains"])
        self.assertEqual(data["bookings"]["deleted_total"], 1)
        self.assertEqual(data["bookings"]["deleted_reasons"][0]["label"], "Неявка на прием")
        self.assertTrue(data["revenue"]["discount_control"]["available"])
        self.assertIn("negative_rows", data["revenue"]["discount_control"])
        self.assertTrue(data["revenue"]["payment_methods"]["available"])
        self.assertEqual(
            {row["method"] for row in data["revenue"]["payment_methods"]["rows"]},
            {"Безналичные", "Наличные"},
        )
        self.assertTrue(data["marketing"]["available"])
        self.assertEqual(data["marketing"]["sources"][0]["source"], "Instagram")

    def test_funnel_days_to_treatment_is_calculated_from_dates(self):
        import nn_bi
        payload = {
            "treatments": [
                {
                    "patient_key": "p1",
                    "first_date": "2026-06-05",
                    "doctors": {"evlk": ["Врач А"], "sclero": [], "mini": []},
                }
            ],
            "repeats": [
                {"patient_key": "p1", "date": "2026-06-10", "doctor": "Врач А"}
            ],
        }
        result = nn_bi._funnel(
            payload,
            {
                "primaries": [
                    {
                        "patient_key": "p1",
                        "date": "2026-06-01",
                        "doctor": "Врач А",
                        "indications": {"evlk": True, "sclero": False, "mini": False},
                    }
                ]
            },
        )
        self.assertEqual(result["days_to_treatment"]["average"], 4.0)
        self.assertEqual(result["days_to_treatment"]["median"], 4)
        self.assertEqual(result["conversion_primary_to_treatment"], 100.0)

    def test_doctor_transition_counts_primary_to_treatment_doctor(self):
        import nn_bi
        rows = nn_bi._doctor_transitions(
            {},
            {
                "primaries": [
                    {"patient_key": "p1", "doctor": "Врач первичного"},
                    {"patient_key": "p2", "doctor": "Врач первичного"},
                ],
                "treatments": [
                    {"patient_key": "p1", "doctors": {"evlk": ["Врач лечения"], "sclero": [], "mini": []}},
                    {"patient_key": "p2", "doctors": {"evlk": [], "sclero": ["Врач лечения"], "mini": []}},
                ],
            },
        )
        self.assertEqual(
            rows,
            [{"primary_doctor": "Врач первичного", "treatment_doctor": "Врач лечения", "patients": 2}],
        )

    def test_relations_support_bar_and_heatmap(self):
        bar = self.client_for(2).get(
            "/api/nn/bi/relations?clinic_id=1&date_from=2026-06-01&date_to=2026-06-30&metric=revenue&dim1=doctor"
        )
        self.assertEqual(bar.status_code, 200)
        b = bar.get_json()
        self.assertEqual(b["mode"], "bar")
        self.assertTrue(b["rows"])

        heat = self.client_for(2).get(
            "/api/nn/bi/relations?clinic_id=1&date_from=2026-06-01&date_to=2026-06-30&metric=patients&dim1=doctor&dim2=weekday"
        )
        self.assertEqual(heat.status_code, 200)
        h = heat.get_json()
        self.assertEqual(h["mode"], "heatmap")
        self.assertTrue(h["x"])
        self.assertTrue(h["y"])
        self.assertEqual(len(h["matrix"]), len(h["y"]))

    def test_bi_export_has_expected_sheets(self):
        response = self.client_for(2).get(
            "/nn/bi/export.xlsx?clinic_id=1&date_from=2026-06-01&date_to=2026-06-30"
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response.mimetype)
        wb = load_workbook(io.BytesIO(response.data), read_only=True, data_only=True)
        try:
            for name in (
                "Обзор", "Динамика", "Воронка", "Врачи", "Переходы врачей",
                "Когорты", "Услуги", "Цепочки услуг", "Запись и потери",
                "Финансы", "Маркетинг",
            ):
                self.assertIn(name, wb.sheetnames)
        finally:
            wb.close()

    def test_upload_only_user_has_no_bi_access(self):
        client = self.client_for(3)
        self.assertEqual(client.get("/nn/bi/").status_code, 403)
        self.assertEqual(client.get("/api/nn/bi/summary?clinic_id=1").status_code, 403)
        self.assertEqual(client.get("/api/nn/bi/relations?clinic_id=1").status_code, 403)
        self.assertEqual(client.get("/nn/bi/export.xlsx?clinic_id=1").status_code, 403)
        self.assertEqual(client.get("/nn/graphics/").status_code, 403)

    def test_optional_payment_parser_extracts_methods_and_period(self):
        rows = nn_normalize._payment_control_rows_from_text(
            """Период: 01.06.2026 00:00 - 20.06.2026 23:59
ВРАЧ А
всего Безнал. 2 373 000,0017 000,00356 000,00
ВРАЧ Б
всего Наличные 2 156 000,00 0,00 156 000,00
""",
            ["Врач А", "Врач Б"],
        )
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["method"], "Безналичные")
        self.assertEqual(rows[0]["doctor"], "Врач А")
        self.assertEqual(rows[0]["period_start"], "2026-06-01")
        self.assertEqual(rows[0]["period_end"], "2026-06-20")
        self.assertEqual(rows[0]["gross_amount"], 373000.0)
        self.assertEqual(rows[0]["discount"], 17000.0)
        self.assertEqual(rows[0]["amount"], 356000.0)
        self.assertEqual(rows[1]["method"], "Наличные")

        kazakh_rows = nn_normalize._payment_control_rows_from_text(
            """Период: 01.06.2026 00:00 - 20.06.2026 23:59
ЫСМАЙЫЛ ДІНМҰХАММЕД ƏБДІРАШИДҰЛЫ
всего Безнал. 36 1 452 000,0017 000,001 435 000,00
""",
            ["Ысмайыл Дінмұхаммед Әбдірашидұлы"],
        )
        self.assertEqual(kazakh_rows[0]["doctor"], "Ысмайыл Дінмұхаммед Әбдірашидұлы")
        self.assertEqual(kazakh_rows[0]["amount"], 1435000.0)

    def test_new_normalization_helpers_preserve_bi_dimensions(self):
        self.assertEqual(nn_normalize._parse_datetime("27.05.2026, 13:47"), "2026-05-27T13:47")
        self.assertEqual(nn_normalize._parse_datetime("2026-05-27T13:47:00"), "2026-05-27T13:47")
        payload = nn_normalize.normalize_rows(
            [{"name": "Иванова Анна", "dob": "1990-01-01", "gender": "Женский", "source": "Instagram", "visits_count": 2, "iin": "", "chart": "10", "phone": "7000000010", "note": "", "source_amount": 0}],
            [{"date": "2026-06-01", "created_at": "2026-05-27T13:47", "start": "10:00", "end": "10:30", "doctor": "Врач А", "patient": "Иванова Анна", "iin": "", "chart": "10", "phone": "7000000010", "repeat_field": "", "note": "", "visit_type": "Амбулаторно", "help_type": "", "appeal_reason": "", "diagnoses": "", "services_text": "Прием хирурга флеболога с УЗИ", "amount": 6000}],
            [], [], [], 0,
        )
        patient = payload["patients"][0]
        self.assertEqual(patient["gender"], "Женский")
        self.assertEqual(patient["source"], "Instagram")
        self.assertEqual(patient["visits_count"], 2)
        self.assertEqual(payload["visits"][0]["created_at"], "2026-05-27T13:47")


if __name__ == "__main__":
    unittest.main()
