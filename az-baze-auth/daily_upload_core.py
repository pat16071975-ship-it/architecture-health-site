import hashlib
import json
import re
import sqlite3
from collections import defaultdict
from datetime import date, datetime
from io import BytesIO

from flask import abort, g, render_template, request
from openpyxl import load_workbook
from werkzeug.exceptions import HTTPException

import ident_import
from app import DB_PATH, audit, csrf_token, db, iso_now, permission_required, require_csrf, user_permissions

UPLOAD_SCHEMA = """
PRAGMA foreign_keys = ON;
CREATE TABLE IF NOT EXISTS daily_uploads (
    data_date TEXT PRIMARY KEY,
    completed_filename TEXT NOT NULL,
    services_filename TEXT NOT NULL,
    completed_sha256 TEXT NOT NULL,
    services_sha256 TEXT NOT NULL,
    normalized_json TEXT NOT NULL,
    revision INTEGER NOT NULL DEFAULT 1,
    uploaded_by INTEGER,
    uploaded_at TEXT NOT NULL,
    FOREIGN KEY (uploaded_by) REFERENCES users(id) ON DELETE SET NULL
);
CREATE TABLE IF NOT EXISTS upload_log (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    data_date TEXT,
    status TEXT NOT NULL,
    details TEXT,
    completed_filename TEXT,
    services_filename TEXT,
    uploaded_by INTEGER,
    created_at TEXT NOT NULL,
    FOREIGN KEY (uploaded_by) REFERENCES users(id) ON DELETE SET NULL
);
CREATE INDEX IF NOT EXISTS idx_upload_log_created_at ON upload_log(created_at DESC);
"""


def _init_schema():
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(DB_PATH)
    try:
        conn.executescript(UPLOAD_SCHEMA)
        conn.commit()
    finally:
        conn.close()


def _plain(value):
    if isinstance(value, defaultdict):
        value = dict(value)
    if isinstance(value, dict):
        return {str(k): _plain(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_plain(v) for v in value]
    return value


def _fmt_cell(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value).replace(".", ",")
    return str(value)


def _decode_text(raw):
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16")
    for enc in ("utf-8-sig", "cp1251"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            pass
    raise ValueError("Не удалось прочитать кодировку файла.")


def _table_candidates(raw, filename):
    lower = (filename or "").lower()
    if raw.startswith(b"PK") or lower.endswith(".xlsx"):
        try:
            wb = load_workbook(BytesIO(raw), read_only=True, data_only=True)
        except Exception as exc:
            raise ValueError("Excel-файл повреждён или имеет неподдерживаемую структуру.") from exc
        candidates = []
        try:
            for ws in wb.worksheets:
                lines = []
                for row in ws.iter_rows(values_only=True):
                    lines.append("\t".join(_fmt_cell(v) for v in row))
                text = "\n".join(lines)
                if text.strip():
                    candidates.append((ws.title, text))
        finally:
            wb.close()
        if not candidates:
            raise ValueError("Excel-файл не содержит данных.")
        return candidates
    return [("текст", _decode_text(raw))]


def _parse_file(file_storage, kind):
    raw = file_storage.read()
    if not raw:
        raise ValueError("Выбран пустой файл.")
    if len(raw) > 25 * 1024 * 1024:
        raise ValueError("Размер файла превышает 25 МБ.")
    candidates = _table_candidates(raw, file_storage.filename or "")
    parser = ident_import._parse_completed if kind == "completed" else ident_import._parse_revenue
    last_error = None
    for sheet_name, text in candidates:
        try:
            return raw, parser(text), sheet_name
        except HTTPException as exc:
            last_error = exc
        except Exception as exc:
            last_error = exc
    label = "«Завершённые приёмы»" if kind == "completed" else "«Выполненные услуги»"
    raise ValueError(f"Не удалось распознать структуру файла {label}.") from last_error


def _sha(raw):
    return hashlib.sha256(raw).hexdigest()


def _load_report(date_value):
    row = db().execute("SELECT payload FROM report_data WHERE date=?", (date_value,)).fetchone()
    if not row:
        return {}
    try:
        value = json.loads(row["payload"])
        return value if isinstance(value, dict) else {}
    except (TypeError, ValueError):
        return {}


def _latest_prior_report(first_date):
    month = first_date[:7]
    row = db().execute(
        "SELECT date,payload FROM report_data WHERE date < ? AND substr(date,1,7)=? ORDER BY date DESC LIMIT 1",
        (first_date, month),
    ).fetchone()
    if not row:
        return {}
    try:
        value = json.loads(row["payload"])
        return value if isinstance(value, dict) else {}
    except (TypeError, ValueError):
        return {}


def _daily_delta(normalized):
    items = normalized.get("items", [])
    overall = normalized.get("overall", {})
    doctors = normalized.get("doctors", {})
    lab_invoices = normalized.get("lab_invoices", [])
    data_date = normalized["data_date"]
    staff = defaultdict(float)
    for row in items:
        staff[row["staff"]] += float(row.get("amount") or 0)
    total_revenue = sum(staff.values())
    lab_revenue = staff.get("Казанцев Л. Е.", 0.0)
    day_counts = overall.get(data_date, {})
    primary = int(day_counts.get("Первичные", 0) or 0)
    repeat = int(day_counts.get("Повторные", 0) or 0) + int(day_counts.get("Отконсультированные", 0) or 0)
    dent_primary = structure_primary = dent_repeat = structure_repeat = 0
    day_doctors = doctors.get(data_date, {})
    for staff_name, value in day_doctors.get("Первичные", {}).items():
        if staff_name in ident_import.DENTISTS:
            dent_primary += int(value or 0)
        elif staff_name in ident_import.STRUCTURE_DOCTORS:
            structure_primary += int(value or 0)
    for kind in ("Повторные", "Отконсультированные"):
        for staff_name, value in day_doctors.get(kind, {}).items():
            if staff_name in ident_import.DENTISTS:
                dent_repeat += int(value or 0)
            elif staff_name in ident_import.STRUCTURE_DOCTORS:
                structure_repeat += int(value or 0)
    lab_orders = len({(str(invoice), str(d)) for invoice, d in lab_invoices if str(d) == data_date})
    return {
        "factMedicine": round(total_revenue - lab_revenue, 2),
        "factLab": round(lab_revenue, 2),
        "primary": primary,
        "repeat": repeat,
        "dentPrimary": dent_primary,
        "dentRepeat": dent_repeat,
        "clinicPrimary": structure_primary,
        "clinicRepeat": structure_repeat,
        "dentists": {full: round(staff.get(short, 0.0), 2) for short, full in ident_import.DENTISTS.items()},
        "clinicDocs": {full: round(staff.get(short, 0.0), 2) for short, full in ident_import.STRUCTURE_DOCTORS.items()},
        "labOrders": lab_orders,
        "labRevenue": round(lab_revenue, 2),
    }


def _active_month_rows(month):
    rows = db().execute(
        "SELECT data_date,normalized_json FROM daily_uploads WHERE substr(data_date,1,7)=? ORDER BY data_date",
        (month,),
    ).fetchall()
    result = []
    for row in rows:
        try:
            result.append((row["data_date"], json.loads(row["normalized_json"])))
        except (TypeError, ValueError):
            continue
    return result


def _rebuild_management(month, source):
    rows = _active_month_rows(month)
    if not rows:
        return {}
    baseline = _latest_prior_report(rows[0][0])
    current = {
        "factMedicine": float(baseline.get("factMedicine") or 0),
        "factLab": float(baseline.get("factLab") or 0),
        "primary": int(baseline.get("primary") or 0),
        "repeat": int(baseline.get("repeat") or 0),
        "dentPrimary": int(baseline.get("dentPrimary") or 0),
        "dentRepeat": int(baseline.get("dentRepeat") or 0),
        "clinicPrimary": int(baseline.get("clinicPrimary") or 0),
        "clinicRepeat": int(baseline.get("clinicRepeat") or 0),
        "dentists": {name: float((baseline.get("dentists") or {}).get(name) or 0) for name in ident_import.DENTISTS.values()},
        "clinicDocs": {name: float((baseline.get("clinicDocs") or {}).get(name) or 0) for name in ident_import.STRUCTURE_DOCTORS.values()},
        "labOrders": int(baseline.get("labOrders") or 0),
        "labRevenue": float(baseline.get("labRevenue") or baseline.get("factLab") or 0),
    }
    rebuilt = {}
    for data_date, normalized in rows:
        delta = _daily_delta(normalized)
        for key in ("factMedicine", "factLab", "primary", "repeat", "dentPrimary", "dentRepeat", "clinicPrimary", "clinicRepeat", "labOrders", "labRevenue"):
            current[key] += delta[key]
        for name, value in delta["dentists"].items():
            current["dentists"][name] = current["dentists"].get(name, 0) + value
        for name, value in delta["clinicDocs"].items():
            current["clinicDocs"][name] = current["clinicDocs"].get(name, 0) + value
        existing = _load_report(data_date)
        rebuilt[data_date] = {
            "date": data_date,
            "plan": existing.get("plan", ""),
            "factMedicine": round(current["factMedicine"], 2),
            "factLab": round(current["factLab"], 2),
            "pp25": existing.get("pp25", ""),
            "avg25": existing.get("avg25", ""),
            "primary": current["primary"],
            "repeat": current["repeat"],
            "dentPrimary": current["dentPrimary"],
            "dentRepeat": current["dentRepeat"],
            "dentists": {k: round(v, 2) for k, v in current["dentists"].items()},
            "clinicPrimary": current["clinicPrimary"],
            "clinicRepeat": current["clinicRepeat"],
            "clinicDocs": {k: round(v, 2) for k, v in current["clinicDocs"].items()},
            "labOrders": current["labOrders"],
            "labRevenue": round(current["labRevenue"], 2),
            "leadsDent": existing.get("leadsDent", ""),
            "leadsDentLost": existing.get("leadsDentLost", ""),
            "leadsClinic": existing.get("leadsClinic", ""),
            "leadsReserve": existing.get("leadsReserve", ""),
            "_source": source,
            "_aggregation": "month_to_date",
        }
    return rebuilt


def _service_through(data):
    for value in (data.get("id"), data.get("source")):
        match = re.search(r"(20\d{2}-\d{2}-\d{2})", str(value or ""))
        if match:
            return match.group(1)
    return None


def _ensure_month(data, year, month):
    if int(data.get("year") or year) != year:
        raise ValueError("Год файла не совпадает с годом базы «Аналитики услуг».")
    label = ident_import.MONTH_NAMES[month - 1]
    months = data.setdefault("months", [])
    if label in months:
        month_index = months.index(label)
    else:
        if len(months) != month - 1:
            raise ValueError("Нельзя добавить этот месяц в «Аналитику услуг»: нарушена последовательность месяцев.")
        months.append(label)
        month_index = len(months) - 1
    directions = data.setdefault("directions", {})
    ident_import._rename_structure_direction(directions)
    for direction in directions.values():
        for doctor in direction.get("doctors", {}).values():
            for series in doctor.values():
                if isinstance(series, list):
                    while len(series) <= month_index:
                        series.append([0, 0])
    return month_index


def _apply_service_items(data, items, year, month, sign):
    month_index = _ensure_month(data, year, month)
    directions = data["directions"]
    for row in items:
        if row["staff"] in ident_import.DENTISTS:
            direction_name, doctor_name = "Стоматология", ident_import.DENTISTS[row["staff"]]
            category = ident_import._classify_dent(row["group"], row["service"])
        elif row["staff"] in ident_import.STRUCTURE_DOCTORS:
            direction_name, doctor_name = "Отделение структуры", ident_import.STRUCTURE_DOCTORS[row["staff"]]
            category = ident_import._classify_structure(row["group"], row["service"])
        elif row["staff"] in ident_import.LAB_DOCTORS:
            direction_name, doctor_name = "Лаборатория", ident_import.LAB_DOCTORS[row["staff"]]
            category = ident_import._classify_lab(row["group"], row["service"])
        else:
            continue
        direction = directions.get(direction_name)
        if not direction:
            continue
        doctor = direction.get("doctors", {}).get(doctor_name)
        if not doctor or category not in direction.get("categories", []):
            continue
        series = doctor.setdefault(category, [])
        while len(series) <= month_index:
            series.append([0, 0])
        series[month_index][0] = max(0, float(series[month_index][0] or 0) + sign * float(row.get("qty") or 0))
        series[month_index][1] = max(0, float(series[month_index][1] or 0) + sign * float(row.get("amount") or 0))
    return month_index


def _save_blob(conn, key, value, now):
    conn.execute(
        "INSERT INTO report_blobs(key,payload,updated_by,updated_at) VALUES(?,?,?,?) ON CONFLICT(key) DO UPDATE SET payload=excluded.payload,updated_by=excluded.updated_by,updated_at=excluded.updated_at",
        (key, json.dumps(value, ensure_ascii=False, separators=(",", ":")), g.user["id"], now),
    )


def _log(conn, status, data_date, details, completed_name, services_name):
    conn.execute(
        "INSERT INTO upload_log(data_date,status,details,completed_filename,services_filename,uploaded_by,created_at) VALUES(?,?,?,?,?,?,?)",
        (data_date, status, details, completed_name, services_name, g.user["id"], iso_now()),
    )


def _process_daily_upload(completed_file, services_file):
    perms = user_permissions(g.user)
    if "upload_completed" not in perms or "upload_services" not in perms:
        abort(403)

    completed_raw, completed_parsed, completed_sheet = _parse_file(completed_file, "completed")
    services_raw, services_parsed, services_sheet = _parse_file(services_file, "services")
    overall, doctors = completed_parsed
    items, lab_invoices, (year, month) = services_parsed

    completed_dates = sorted(overall)
    service_dates = sorted({row["date"] for row in items})
    if len(completed_dates) != 1:
        raise ValueError("Файл «Завершённые приёмы» должен содержать данные ровно за один день.")
    if len(service_dates) != 1:
        raise ValueError("Файл «Выполненные услуги» должен содержать данные ровно за один день.")
    if completed_dates[0] != service_dates[0]:
        raise ValueError("Даты в двух файлах не совпадают.")

    data_date = completed_dates[0]
    if data_date[:7] != f"{year:04d}-{month:02d}":
        raise ValueError("Месяц в двух файлах не совпадает.")

    completed_hash, services_hash = _sha(completed_raw), _sha(services_raw)
    completed_name = completed_file.filename or "Завершённые приёмы"
    services_name = services_file.filename or "Выполненные услуги"
    normalized = {
        "data_date": data_date,
        "items": _plain(items),
        "lab_invoices": _plain(lab_invoices),
        "overall": _plain(overall),
        "doctors": _plain(doctors),
    }

    conn = db()
    existing = conn.execute("SELECT * FROM daily_uploads WHERE data_date=?", (data_date,)).fetchone()
    if existing and existing["completed_sha256"] == completed_hash and existing["services_sha256"] == services_hash:
        _log(conn, "duplicate", data_date, "Повторная загрузка идентичного комплекта; данные не изменялись.", completed_name, services_name)
        conn.commit()
        return {"status": "duplicate", "message": f"Данные за {data_date} уже загружены. Повторно ничего не изменено.", "data_date": data_date}

    replacing = bool(existing)
    if replacing and "upload_replace" not in perms:
        _log(conn, "denied_replace", data_date, "Комплект за дату уже существует; нет права на замену.", completed_name, services_name)
        conn.commit()
        raise ValueError("За эту дату данные уже загружены. Для замены требуется отдельное право.")

    service_data = ident_import._load_blob("az-service-analytics-v1")
    if not service_data:
        raise ValueError("База «Аналитики услуг» ещё не подготовлена.")
    if not replacing:
        through = _service_through(service_data)
        if through and data_date <= through:
            raise ValueError(f"Дата {data_date} уже входит в ранее загруженную сводную базу по {through}. Ежедневную загрузку начинайте со следующего дня.")

    old_normalized = None
    if replacing:
        try:
            old_normalized = json.loads(existing["normalized_json"])
        except (TypeError, ValueError):
            raise ValueError("Нельзя безопасно заменить прошлую загрузку: сохранённая версия повреждена.")

    if old_normalized:
        _apply_service_items(service_data, old_normalized.get("items", []), year, month, -1)
    month_index = _apply_service_items(service_data, items, year, month, +1)
    old_through = _service_through(service_data) or data_date
    last_known = max(data_date, old_through)
    service_data["id"] = f"az-services-{year}-through-{last_known}-v1"
    service_data["source"] = f"Ежедневные загрузки AZ-BAZE through {last_known}"
    finance_blobs = ident_import._pad_financial_months(month_index)

    now = iso_now()
    source = f"AZ-BAZE daily: {services_name} + {completed_name}"
    conn.execute("BEGIN")
    try:
        if replacing:
            revision = int(existing["revision"] or 1) + 1
            conn.execute(
                "UPDATE daily_uploads SET completed_filename=?,services_filename=?,completed_sha256=?,services_sha256=?,normalized_json=?,revision=?,uploaded_by=?,uploaded_at=? WHERE data_date=?",
                (completed_name, services_name, completed_hash, services_hash, json.dumps(normalized, ensure_ascii=False, separators=(",", ":")), revision, g.user["id"], now, data_date),
            )
        else:
            conn.execute(
                "INSERT INTO daily_uploads(data_date,completed_filename,services_filename,completed_sha256,services_sha256,normalized_json,revision,uploaded_by,uploaded_at) VALUES(?,?,?,?,?,?,1,?,?)",
                (data_date, completed_name, services_name, completed_hash, services_hash, json.dumps(normalized, ensure_ascii=False, separators=(",", ":")), g.user["id"], now),
            )

        management = _rebuild_management(data_date[:7], source)
        for day, record in management.items():
            conn.execute(
                "INSERT INTO report_data(date,payload,updated_by,updated_at) VALUES(?,?,?,?) ON CONFLICT(date) DO UPDATE SET payload=excluded.payload,updated_by=excluded.updated_by,updated_at=excluded.updated_at",
                (day, json.dumps(record, ensure_ascii=False, separators=(",", ":")), g.user["id"], now),
            )

        _save_blob(conn, "az-service-analytics-v1", service_data, now)
        for key, value in finance_blobs.items():
            _save_blob(conn, key, value, now)
        _log(conn, "replaced" if replacing else "imported", data_date, f"Завершённые приёмы: лист {completed_sheet}; Выполненные услуги: лист {services_sheet}; management={len(management)}", completed_name, services_name)
        conn.commit()
    except Exception:
        conn.rollback()
        raise

    audit(
        "daily_reports_replaced" if replacing else "daily_reports_imported",
        target_user_id=g.user["id"],
        details=f"date={data_date}; completed={completed_name}; services={services_name}",
    )
    return {
        "status": "replaced" if replacing else "imported",
        "message": f"Готово. Данные за {data_date} {'заменены' if replacing else 'загружены'} и разнесены по управленческому отчёту и «Аналитике услуг».",
        "data_date": data_date,
    }


def _history():
    return db().execute(
        "SELECT l.*,u.full_name AS user_name FROM upload_log l LEFT JOIN users u ON u.id=l.uploaded_by ORDER BY l.id DESC LIMIT 60"
    ).fetchall()


def register_daily_upload(app):
    _init_schema()

    @app.route("/uploads/", methods=["GET", "POST"])
    @permission_required("section5")
    def uploads_page():
        perms = user_permissions(g.user)
        result = None
        error = None
        if request.method == "POST":
            require_csrf()
            completed = request.files.get("completed")
            services = request.files.get("services")
            if not completed or not completed.filename:
                error = "Выберите файл «Завершённые приёмы»."
            elif not services or not services.filename:
                error = "Выберите файл «Выполненные услуги»."
            else:
                try:
                    result = _process_daily_upload(completed, services)
                except ValueError as exc:
                    error = str(exc)

        latest = db().execute(
            "SELECT data_date,revision,uploaded_at FROM daily_uploads ORDER BY data_date DESC LIMIT 1"
        ).fetchone()
        return render_template(
            "uploads.html",
            csrf=csrf_token(),
            perms=perms,
            result=result,
            error=error,
            latest=latest,
            history=_history() if "upload_history" in perms else [],
            can_daily=("upload_completed" in perms and "upload_services" in perms),
            can_replace=("upload_replace" in perms),
        )
