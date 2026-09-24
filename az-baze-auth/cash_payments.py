import hashlib
import json
import re
from collections import defaultdict
from datetime import datetime
from io import BytesIO

from openpyxl import load_workbook


OOO_KKM = "0531380019042729"
IP_KKM = "0463880019042725"
LEGAL_BY_KKM = {OOO_KKM: "ooo", IP_KKM: "ip"}
_PROVIDER_MAPS = {"dent": {}, "structure": {}, "lab": {}}


def configure_providers(dentists, structure_doctors, lab_doctors):
    _PROVIDER_MAPS["dent"] = dict(dentists or {})
    _PROVIDER_MAPS["structure"] = dict(structure_doctors or {})
    _PROVIDER_MAPS["lab"] = dict(lab_doctors or {})

MONTHS_RU = {
    "янв": 1, "фев": 2, "мар": 3, "апр": 4, "май": 5, "июн": 6,
    "июл": 7, "авг": 8, "сен": 9, "сент": 9, "окт": 10, "ноя": 11, "дек": 12,
}
REQUIRED_HEADERS = {
    "date": "Дата и время",
    "operation": "Операция",
    "due": "Сумма к оплате (₽)",
    "movement": "Движение ДС (₽)",
    "kkm": "ККМ",
}
INVOICE_OPERATION_RE = re.compile(r"^№\d+")
DEBT_OPERATION_RE = re.compile(r"^Задолженность по счету №\d+", re.IGNORECASE)


CASH_SCHEMA = """
CREATE TABLE IF NOT EXISTS cash_receipts_daily (
    data_date TEXT PRIMARY KEY,
    payload_json TEXT NOT NULL,
    source_filename TEXT NOT NULL,
    source_sha256 TEXT NOT NULL,
    imported_by INTEGER,
    imported_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_cash_receipts_daily_date ON cash_receipts_daily(data_date);
"""


def init_schema(conn):
    conn.executescript(CASH_SCHEMA)


def sha256(raw):
    return hashlib.sha256(raw).hexdigest()


def _money(value):
    text = str(value or "").replace("₽", "").replace("\xa0", " ").strip()
    text = text.replace(" ", "").replace(",", ".")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _date(value):
    text = str(value or "").strip().lower().replace("ё", "е")
    if not text:
        return None
    first = text.splitlines()[0].strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(first, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    match = re.match(r"^(\d{1,2})\s+([а-я]+)\s+(20\d{2})$", first)
    if not match:
        return None
    month = MONTHS_RU.get(match.group(2)[:4]) or MONTHS_RU.get(match.group(2)[:3])
    if not month:
        return None
    return f"{int(match.group(3)):04d}-{month:02d}-{int(match.group(1)):02d}"


def _blank_day():
    return {
        "billedTotal": 0.0,
        "cashOOO": 0.0,
        "cashIP": 0.0,
        "cashTotal": 0.0,
        "cashUnallocated": 0.0,
        "factMedicine": 0.0,
        "factLab": 0.0,
        "dentists": {},
        "clinicDocs": {},
        "labRevenue": 0.0,
        "dentistsLegal": {},
        "clinicDocsLegal": {},
        "labLegal": {"ooo": 0.0, "ip": 0.0},
        "dentCashOOO": 0.0,
        "dentCashIP": 0.0,
        "clinicCashOOO": 0.0,
        "clinicCashIP": 0.0,
        "labCashOOO": 0.0,
        "labCashIP": 0.0,
    }


def _add_legal(container, name, legal, value):
    target = container.setdefault(name, {"ooo": 0.0, "ip": 0.0})
    target[legal] = float(target.get(legal) or 0) + value


def _find_provider(operation):
    candidates = []
    for direction, mapping in _PROVIDER_MAPS.items():
        for short, full in mapping.items():
            candidates.append((short, direction, full))
    for short, direction, full in sorted(candidates, key=lambda item: len(item[0]), reverse=True):
        if short and short in operation:
            return direction, full
    return None, None


def _parse_rows(rows):
    rows = list(rows)
    if not rows:
        raise ValueError("Файл «Счета и оплаты» пуст.")
    headers = [str(part or "").strip() for part in rows[0]]
    index = {}
    for key, label in REQUIRED_HEADERS.items():
        try:
            index[key] = headers.index(label)
        except ValueError as exc:
            raise ValueError(f"В файле «Счета и оплаты» нет обязательной колонки «{label}».") from exc

    result = defaultdict(_blank_day)
    recognized = 0
    for source_row in rows[1:]:
        cols = list(source_row)
        if not cols or str(cols[0] or "").strip().lower() == "итого":
            continue
        need = max(index.values()) + 1
        if len(cols) < need:
            cols += [""] * (need - len(cols))
        data_date = _date(cols[index["date"]])
        if not data_date:
            continue
        operation = str(cols[index["operation"]] or "").strip()
        due = _money(cols[index["due"]])
        movement = _money(cols[index["movement"]])
        kkm = re.sub(r"\D", "", str(cols[index["kkm"]] or ""))
        day = result[data_date]
        recognized += 1

        if due is not None and due > 0 and (
            INVOICE_OPERATION_RE.match(operation) or DEBT_OPERATION_RE.match(operation)
        ):
            day["billedTotal"] += due

        if movement is not None and movement > 0 and kkm and kkm not in LEGAL_BY_KKM:
            raise ValueError(f"В файле «Счета и оплаты» обнаружена неизвестная ККМ: {kkm}.")

        legal = LEGAL_BY_KKM.get(kkm)
        if legal is None or movement is None or movement <= 0:
            continue
        if legal == "ooo":
            day["cashOOO"] += movement
        else:
            day["cashIP"] += movement
        day["cashTotal"] += movement

        direction, provider = _find_provider(operation) if operation.startswith("№") else (None, None)
        if direction == "dent":
            day["dentists"][provider] = float(day["dentists"].get(provider) or 0) + movement
            _add_legal(day["dentistsLegal"], provider, legal, movement)
            day["factMedicine"] += movement
            day["dentCashOOO" if legal == "ooo" else "dentCashIP"] += movement
        elif direction == "structure":
            day["clinicDocs"][provider] = float(day["clinicDocs"].get(provider) or 0) + movement
            _add_legal(day["clinicDocsLegal"], provider, legal, movement)
            day["factMedicine"] += movement
            day["clinicCashOOO" if legal == "ooo" else "clinicCashIP"] += movement
        elif direction == "lab":
            day["factLab"] += movement
            day["labRevenue"] += movement
            day["labLegal"][legal] += movement
            day["labCashOOO" if legal == "ooo" else "labCashIP"] += movement
        else:
            day["cashUnallocated"] += movement

    if not recognized or not result:
        raise ValueError("Файл «Счета и оплаты» не содержит распознаваемых операций.")

    normalized = {}
    for data_date, day in result.items():
        allocated = float(day["factMedicine"]) + float(day["factLab"])
        day["cashUnallocated"] = max(0.0, float(day["cashTotal"]) - allocated)
        normalized[data_date] = _round_payload(day)
    return dict(sorted(normalized.items()))


def parse_file(raw, filename=""):
    if not raw:
        raise ValueError("Файл «Счета и оплаты» пуст.")
    if len(raw) > 40 * 1024 * 1024:
        raise ValueError("Размер файла «Счета и оплаты» превышает 40 МБ.")
    lower = str(filename or "").lower()
    if raw.startswith(b"PK") or lower.endswith(".xlsx"):
        try:
            workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
        except Exception as exc:
            raise ValueError("Excel-файл «Счета и оплаты» повреждён или имеет неподдерживаемую структуру.") from exc
        last_error = None
        try:
            for worksheet in workbook.worksheets:
                try:
                    parsed = _parse_rows(worksheet.iter_rows(values_only=True))
                    return parsed, worksheet.title
                except ValueError as exc:
                    last_error = exc
        finally:
            workbook.close()
        if last_error:
            raise last_error
        raise ValueError("Excel-файл «Счета и оплаты» не содержит данных.")

    text = raw.decode("utf-8-sig")
    rows = [line.split("\t") for line in text.splitlines()]
    return _parse_rows(rows), "текст"


def _round_payload(value):
    if isinstance(value, dict):
        return {str(k): _round_payload(v) for k, v in value.items()}
    if isinstance(value, float):
        return round(value, 2)
    return value


def replace_range(conn, daily, source_filename, source_sha256, imported_by, imported_at):
    if not daily:
        raise ValueError("Файл «Счета и оплаты» не содержит данных для сохранения.")
    dates = sorted(daily)
    # Replace only dates present in the uploaded file. This keeps prior daily
    # receipts when administrators upload the next day or a partial period.
    for data_date in dates:
        conn.execute(
            "DELETE FROM cash_receipts_daily WHERE data_date=?",
            (data_date,),
        )
        conn.execute(
            "INSERT INTO cash_receipts_daily(data_date,payload_json,source_filename,source_sha256,imported_by,imported_at) VALUES(?,?,?,?,?,?)",
            (
                data_date,
                json.dumps(daily[data_date], ensure_ascii=False, separators=(",", ":")),
                source_filename,
                source_sha256,
                imported_by,
                imported_at,
            ),
        )
    return sorted({data_date[:7] for data_date in dates})


def _merge_numbers(target, source):
    for key in (
        "billedTotal", "cashOOO", "cashIP", "cashTotal", "cashUnallocated",
        "factMedicine", "factLab", "labRevenue", "dentCashOOO", "dentCashIP",
        "clinicCashOOO", "clinicCashIP", "labCashOOO", "labCashIP",
    ):
        target[key] = float(target.get(key) or 0) + float(source.get(key) or 0)
    for key in ("dentists", "clinicDocs"):
        for name, amount in (source.get(key) or {}).items():
            target[key][name] = float(target[key].get(name) or 0) + float(amount or 0)
    for key in ("dentistsLegal", "clinicDocsLegal"):
        for name, split in (source.get(key) or {}).items():
            _add_legal(target[key], name, "ooo", float((split or {}).get("ooo") or 0))
            _add_legal(target[key], name, "ip", float((split or {}).get("ip") or 0))
    target["labLegal"]["ooo"] += float((source.get("labLegal") or {}).get("ooo") or 0)
    target["labLegal"]["ip"] += float((source.get("labLegal") or {}).get("ip") or 0)


def month_snapshots(conn, month):
    rows = conn.execute(
        "SELECT data_date,payload_json FROM cash_receipts_daily WHERE substr(data_date,1,7)=? ORDER BY data_date",
        (month,),
    ).fetchall()
    current = _blank_day()
    snapshots = {}
    for row in rows:
        try:
            payload = json.loads(row["payload_json"])
        except (TypeError, ValueError):
            continue
        _merge_numbers(current, payload if isinstance(payload, dict) else {})
        snapshots[str(row["data_date"])] = _round_payload(current)
    return snapshots


def _latest_snapshot(snapshots, data_date):
    eligible = [date for date in snapshots if date <= data_date]
    if not eligible:
        return None
    return snapshots[max(eligible)]


def apply_to_record(record, snapshot, source_label="Счета и оплаты"):
    if not snapshot:
        return record
    result = dict(record or {})
    result.setdefault("billedMedicine", result.get("factMedicine", 0))
    result.setdefault("billedLab", result.get("factLab", result.get("labRevenue", 0)))
    result.setdefault("billedDentists", dict(result.get("dentists") or {}))
    result.setdefault("billedClinicDocs", dict(result.get("clinicDocs") or {}))
    result.setdefault("billedLabRevenue", result.get("labRevenue", result.get("factLab", 0)))

    for key in (
        "billedTotal", "cashOOO", "cashIP", "cashTotal", "cashUnallocated",
        "factMedicine", "factLab", "labRevenue", "dentCashOOO", "dentCashIP",
        "clinicCashOOO", "clinicCashIP", "labCashOOO", "labCashIP",
    ):
        result[key] = snapshot.get(key, 0)
    result["dentists"] = dict(snapshot.get("dentists") or {})
    result["clinicDocs"] = dict(snapshot.get("clinicDocs") or {})
    result["dentistsLegal"] = dict(snapshot.get("dentistsLegal") or {})
    result["clinicDocsLegal"] = dict(snapshot.get("clinicDocsLegal") or {})
    result["labLegal"] = dict(snapshot.get("labLegal") or {"ooo": 0, "ip": 0})
    result["_cash_source"] = source_label
    result["_cash_rule"] = "positive-receipts-only-v1"
    return result


def overlay_record_map(conn, month, records):
    snapshots = month_snapshots(conn, month)
    if not snapshots:
        return records
    for data_date in list(records):
        snapshot = _latest_snapshot(snapshots, data_date)
        if snapshot:
            records[data_date] = apply_to_record(records[data_date], snapshot)
    return records


def overlay_stored_month(conn, month, updated_by, updated_at):
    snapshots = month_snapshots(conn, month)
    if not snapshots:
        return 0
    rows = conn.execute(
        "SELECT date,payload FROM report_data WHERE substr(date,1,7)=? ORDER BY date",
        (month,),
    ).fetchall()
    if not rows:
        # Cash is authoritative even if an old backup contains no management
        # row for this month. Recreate a minimal cash-only month-end record so
        # Fact cannot disappear while cash_receipts_daily still exists.
        last_cash_date = max(snapshots)
        updated = apply_to_record({"date": last_cash_date}, snapshots[last_cash_date])
        updated["date"] = last_cash_date
        conn.execute(
            "INSERT INTO report_data(date,payload,updated_by,updated_at) VALUES(?,?,?,?)",
            (last_cash_date, json.dumps(updated, ensure_ascii=False, separators=(",", ":")), updated_by, updated_at),
        )
        return 1

    parsed = []
    for row in rows:
        try:
            payload = json.loads(row["payload"])
        except (TypeError, ValueError):
            continue
        if isinstance(payload, dict):
            parsed.append((str(row["date"]), payload))
    if not parsed:
        return 0

    last_cash_date = max(snapshots)
    last_report_date, last_report = parsed[-1]
    if last_cash_date > last_report_date:
        clone = dict(last_report)
        clone["date"] = last_cash_date
        parsed.append((last_cash_date, clone))

    changed = 0
    seen = set()
    for data_date, payload in parsed:
        if data_date in seen:
            continue
        seen.add(data_date)
        snapshot = _latest_snapshot(snapshots, data_date)
        if not snapshot:
            continue
        updated = apply_to_record(payload, snapshot)
        updated["date"] = data_date
        conn.execute(
            "INSERT INTO report_data(date,payload,updated_by,updated_at) VALUES(?,?,?,?) "
            "ON CONFLICT(date) DO UPDATE SET payload=excluded.payload,updated_by=excluded.updated_by,updated_at=excluded.updated_at",
            (data_date, json.dumps(updated, ensure_ascii=False, separators=(",", ":")), updated_by, updated_at),
        )
        changed += 1
    return changed
