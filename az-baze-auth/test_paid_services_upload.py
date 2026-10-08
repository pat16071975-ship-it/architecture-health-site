import io
import sqlite3
import unittest

from openpyxl import Workbook
from openpyxl.styles import PatternFill

import ident_import
import paid_services_upload
import upload_reconcile


class PaidServicesUploadTests(unittest.TestCase):
    def workbook_bytes(self, provider_rows, *, total_opening, total_billed, total_paid, total_closing, second_month=False):
        wb = Workbook()
        ws = wb.active
        ws.title = "Выручка по направлениям"
        ws.append([
            "Группа услуг", "Услуги", "", "",
            " Задолженность на нач. периода ", "Сумма со скидкой",
            " Оплачено ", " Задолженность на конец периода ",
        ])
        section_fill = PatternFill("solid", fgColor="D0EAFF")
        provider_fill = PatternFill("solid", fgColor="EAF6FF")

        ws.append(["Услуги", "", "", "", total_opening, total_billed, total_paid, total_closing])
        ws["A2"].fill = section_fill

        invoice = 71000
        for idx, row in enumerate(provider_rows, start=1):
            name, opening, billed, paid, closing = row
            ws.append([name, "", "", "", opening, billed, paid, closing])
            ws.cell(ws.max_row, 1).fill = provider_fill
            ws.append([f"Пациент {idx}", "", "", "", opening, billed, paid, closing])
            date_text = "01.10.2026" if second_month and idx == len(provider_rows) else "16.09.2026"
            ws.append([f"Счет №{invoice} от {date_text} 12:00:00", "", "", "", "-", billed, paid, closing])
            ws.append(["Терапия", f"{date_text} 12:00", f"Услуга {idx}", 1, "-", billed, paid, closing])
            invoice += 1

        stream = io.BytesIO()
        wb.save(stream)
        wb.close()
        return stream.getvalue()

    def test_new_mis_report_parses_paid_and_debt_balances(self):
        raw = self.workbook_bytes(
            [
                ("Иванова А. В.", 0, 300000, 225000, 75000),
                ("Босак Я. С.", 1076700, 313100, 556800, 833000),
                ("Казанцев Л. Е.", 2262210, 573640, 173240, 2662610),
            ],
            total_opening=3338910,
            total_billed=1186740,
            total_paid=955040,
            total_closing=3570610,
        )
        payload, sheet = paid_services_upload.parse_file(raw, "paid-services.xlsx")
        self.assertEqual(sheet, "Выручка по направлениям")
        self.assertEqual(payload["month"], "2026-09")
        self.assertEqual(payload["period"], {"from": "2026-09-16", "to": "2026-09-16"})
        self.assertEqual(payload["totals"]["paid"], 955040)
        self.assertEqual(payload["providers"]["Иванова А. В."]["paid"], 225000)
        self.assertEqual(payload["providers"]["Босак Я. С."]["closing"], 833000)
        self.assertEqual(payload["providers"]["Казанцев Л. Е."]["billed"], 573640)
        self.assertEqual(len(payload["services"]), 3)

    def test_provider_projection_keeps_paid_by_direction_without_ooo_ip(self):
        conn = sqlite3.connect(":memory:")
        conn.row_factory = sqlite3.Row
        upload_reconcile.init_schema(conn)
        upload_reconcile.seed_defaults(conn, ident_import)
        upload_reconcile.resolve_providers(
            conn,
            {
                "Иванова А. В.": {"direction": "dent", "display_name": "Иванова А. В."},
                "Черевко А. С.": {"direction": "structure", "display_name": "Черевко А. С."},
            },
            1,
        )
        conn.commit()

        payload = {
            "providers": {
                "Иванова А. В.": {"opening": 0, "billed": 300000, "paid": 225000, "closing": 75000},
                "Черевко А. С.": {"opening": 16000, "billed": 2500, "paid": 2500, "closing": 16000},
                "Казанцев Л. Е.": {"opening": 10, "billed": 100, "paid": 90, "closing": 20},
                "Новый Врач А. А.": {"opening": 0, "billed": 7000, "paid": 7000, "closing": 0},
            }
        }
        projection = paid_services_upload._provider_projection(payload, conn)
        self.assertEqual(projection["directions"]["dent"]["paid"], 225000)
        self.assertEqual(projection["directions"]["structure"]["paid"], 2500)
        self.assertEqual(projection["directions"]["lab"]["paid"], 90)
        self.assertEqual(projection["doctors"]["dent"]["Иванова А. В."], 225000)
        self.assertEqual(projection["unknown"], ["Новый Врач А. А."])
        conn.close()

    def test_balance_mismatch_is_rejected(self):
        raw = self.workbook_bytes(
            [("Иванова А. В.", 0, 300000, 225000, 75000)],
            total_opening=0,
            total_billed=300000,
            total_paid=225000,
            total_closing=70000,
        )
        with self.assertRaisesRegex(ValueError, "не сходится"):
            paid_services_upload.parse_file(raw, "paid-services.xlsx")

    def test_two_calendar_months_are_rejected(self):
        raw = self.workbook_bytes(
            [
                ("Иванова А. В.", 0, 300000, 225000, 75000),
                ("Босак Я. С.", 0, 100000, 100000, 0),
            ],
            total_opening=0,
            total_billed=400000,
            total_paid=325000,
            total_closing=75000,
            second_month=True,
        )
        with self.assertRaisesRegex(ValueError, "один календарный месяц"):
            paid_services_upload.parse_file(raw, "paid-services.xlsx")


if __name__ == "__main__":
    unittest.main()
