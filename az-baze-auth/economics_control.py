import hashlib
import json
import re
from pathlib import Path

from flask import Response, g, jsonify, request

from app import SITE_ROOT, admin_required, audit, csrf_token, iso_now, permission_required, require_csrf

DATA_PATH = Path("/var/lib/az-baze/economics-control.json")
MONTHS = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь"]
OPU_COLS = [21, 23, 25, 27, 29, 31]
PAYROLL_COLS = [14, 15, 16, 17, 18, 19]


def _text(v):
    return re.sub(r"\s+", " ", str(v or "")).strip()


def _num(v):
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


def _norm_name(v):
    s = _text(v).replace("(с)", "").replace("(c)", "").strip()
    return re.sub(r"\s+", " ", s)


def _employee_id(name):
    return hashlib.sha256(_norm_name(name).lower().encode("utf-8")).hexdigest()[:16]


def _series(ws, row, cols):
    return [_num(ws.cell(row, c).value) for c in cols]


def _find_row(ws, label, start=1):
    target = _text(label).lower()
    for r in range(start, ws.max_row + 1):
        if _text(ws.cell(r, 4).value).lower() == target:
            return r
    return None


def _parse_opu(ws):
    rows = {
        "revenue": _find_row(ws, "ВЫРУЧКА со скидкой, руб.", 38),
        "cost": _find_row(ws, "СЕБЕСТОИМОСТЬ, руб.", 79),
        "mp1": _find_row(ws, "МАРЖИНАЛЬНАЯ ПРИБЫЛЬ 1", 100),
        "operating_profit": _find_row(ws, "ОПЕРАЦИОННАЯ ПРИБЫЛЬ", 100),
        "net_profit": _find_row(ws, "ЧИСТАЯ ПРИБЫЛЬ", 100),
        "materials": _find_row(ws, "Материалы", 79),
        "assistants": _find_row(ws, "Ассистенты стоматологов", 79),
        "hygiene_fot": _find_row(ws, "Гигиенист", 79),
    }
    out = {key: (_series(ws, row, OPU_COLS) if row else [0] * 6) for key, row in rows.items()}
    out["operating_margin"] = [out["operating_profit"][i] / out["revenue"][i] if out["revenue"][i] else None for i in range(6)]
    out["net_margin"] = [out["net_profit"][i] / out["revenue"][i] if out["revenue"][i] else None for i in range(6)]
    out["rows_found"] = rows
    return out


def _group_for_row(row):
    if row <= 44:
        return "Врачи"
    if row <= 90:
        return "Вспомогательный персонал"
    if row <= 103:
        return "АУП"
    if row <= 113:
        return "Администраторы"
    if row <= 118:
        return "Лаборатория"
    return "Штатный маркетинг"


def _function_for(role, group):
    r = role.lower()
    if group == "АУП":
        if "главн" in r and "врач" in r: return "Главный врач"
        if "hr" in r or "кадр" in r: return "HR / кадры"
        if "бухгалтер" in r or "платеж" in r: return "Бухгалтерия / финансы"
        if "маркет" in r or "smm" in r: return "Маркетинг / SMM"
        if "ассистент" in r: return "Бизнес-ассистенты"
        return "Руководство / управление"
    if group == "Администраторы": return "Администраторы / клиентский сервис"
    if group == "Лаборатория": return "Зуботехническая лаборатория"
    if group == "Штатный маркетинг": return "Маркетинг / SMM"
    if group == "Вспомогательный персонал":
        if "ассист" in r or "стаж" in r: return "Ассистенты врачей"
        if "цсо" in r: return "ЦСО"
        if "убор" in r: return "Уборка"
        if "рентген" in r or "diers" in r or "иггт" in r: return "Диагностика"
        if "мед" in r or "сестр" in r: return "Медсёстры"
        return "Вспомогательный персонал"
    return role or group


def _parse_payroll(ws):
    records, seen = [], set()
    subtotal_names = {"врачебный персонал", "отделение терапии", "отделение структуры", "отделение ортопедии", "отделение хирургии", "миофункциональные терапевты", "средний и младший персонал", "цсо", "административный персонал", "администраторы", "зуботехническая лаборатория", "маркетинговый блок", "рентгенолаборант", "оператор иггт", "ассистент стоматолога", "уборщик служебных помещений"}
    for r in range(6, ws.max_row + 1):
        name, role = _norm_name(ws.cell(r, 5).value), _text(ws.cell(r, 4).value)
        if not name or not role or name.lower() in subtotal_names or name.lower() == role.lower():
            continue
        values = [_num(ws.cell(r, c).value) for c in PAYROLL_COLS]
        if not any(abs(v) > 0.0001 for v in values):
            continue
        eid, group = _employee_id(name), _group_for_row(r)
        function = _function_for(role, group)
        key = (eid, role, tuple(round(x, 2) for x in values))
        if key in seen:
            continue
        seen.add(key)
        records.append({"employee_id": eid, "name": name, "role": role, "group": group, "function": function, "department": _text(ws.cell(r, 3).value) or function, "months": values, "total": sum(values), "source_row": r})
    group_names = ["Врачи", "АУП", "Администраторы", "Вспомогательный персонал", "Лаборатория", "Штатный маркетинг"]
    groups = {}
    for gname in group_names:
        recs = [r for r in records if r["group"] == gname]
        monthly = [sum(r["months"][i] for r in recs) for i in range(6)]
        functions = {}
        for rec in recs:
            functions.setdefault(rec["function"], []).append(rec)
        groups[gname] = {"months": monthly, "total": sum(monthly), "functions": functions}
    total_months = [sum(r["months"][i] for r in records) for i in range(6)]
    return {"records": records, "groups": groups, "months": total_months, "total": sum(total_months)}


def _build_payload(wb):
    if not {"ОПУ", "Свод по ЗП new"}.issubset(set(wb.sheetnames)):
        raise ValueError("В файле нет листов ОПУ и/или Свод по ЗП new")
    opu, payroll = _parse_opu(wb["ОПУ"]), _parse_payroll(wb["Свод по ЗП new"])
    revenue = opu["revenue"]
    fot_ratio = [payroll["months"][i] / revenue[i] if revenue[i] else None for i in range(6)]
    hygiene_quality = ["complete" if opu["hygiene_fot"][i] > 0 else "missing_fot" for i in range(6)]
    checks = {"materials_jan_jun": sum(opu["materials"]), "assistants_jan_jun": sum(opu["assistants"]), "payroll_jan_jun": payroll["total"], "payroll_groups_sum": sum(v["total"] for v in payroll["groups"].values())}
    checks["payroll_balanced"] = abs(checks["payroll_jan_jun"] - checks["payroll_groups_sum"]) < 1
    return {"version": 2, "generated_at": iso_now(), "source": "ОТЧЕТНОСТЬ_Эксперт.xlsx", "months": MONTHS, "opu": opu, "payroll": payroll, "fot_revenue_ratio": fot_ratio, "quality": {"Гигиена": hygiene_quality}, "checks": checks, "notes": {"distributed_profit": "Распределённая прибыль не является фактической прибылью врача.", "marketing": "Стоимость первичного визита по общим расходам маркетинга не является CAC без сквозной атрибуции.", "target_margin": "Цель 20% требует выбора: операционная или чистая рентабельность."}}


def _load_data():
    if not DATA_PATH.exists():
        return {"version": 2, "available": False}
    data = json.loads(DATA_PATH.read_text(encoding="utf-8"))
    data["available"] = True
    return data


def register_economics_control(app):
    @app.get("/reports/dashboard.html")
    @permission_required("reports")
    def economics_dashboard_page():
        return Response((SITE_ROOT / "reports" / "dashboard.html").read_text(encoding="utf-8"), mimetype="text/html")

    @app.route("/reports/economics-import/", methods=["GET", "POST"])
    @admin_required
    def economics_import():
        message, error = "", ""
        if request.method == "POST":
            try:
                require_csrf()
                from openpyxl import load_workbook
                f = request.files.get("file")
                if not f or not f.filename.lower().endswith(".xlsx"):
                    raise ValueError("Выберите исходный файл .xlsx")
                payload = _build_payload(load_workbook(f.stream, data_only=True, read_only=True))
                DATA_PATH.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
                DATA_PATH.chmod(0o640)
                audit("economics_source_imported", target_user_id=g.user["id"], details=f"file={f.filename}; payroll={payload['payroll']['total']:.2f}")
                message = "Экономические данные пересчитаны из первичного файла."
            except Exception as exc:
                error = str(exc)
        html = (SITE_ROOT / "reports" / "economics-import.html").read_text(encoding="utf-8")
        return Response(html.replace("{{CSRF}}", csrf_token()).replace("{{MESSAGE}}", message).replace("{{ERROR}}", error), mimetype="text/html")

    @app.get("/api/reports/economics-control")
    @permission_required("reports")
    def economics_control_data():
        try:
            return jsonify(data=_load_data(), error="", csrf=csrf_token())
        except Exception as exc:
            return jsonify(data={"version": 2, "available": False}, error=str(exc), csrf=csrf_token())
