import hashlib
import json
import re
from datetime import date, datetime
from io import BytesIO

from flask import abort, g, jsonify, request
from openpyxl import load_workbook

import ident_import
import upload_reconcile
from app import audit, db, iso_now, permission_required, require_csrf, user_permissions


MAX_FILE_SIZE = 25 * 1024 * 1024
INVOICE_RE = re.compile(r"^Счет\s+№([^\s]+)\s+от\s+(\d{2}\.\d{2}\.\d{4})", re.IGNORECASE)
EXPECTED_HEADERS = {
    "Задолженность на нач. периода",
    "Сумма со скидкой",
    "Оплачено",
    "Задолженность на конец периода",
}


def init_schema(conn):
    conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS paid_services_months (
            month TEXT PRIMARY KEY,
            payload_json TEXT NOT NULL,
            source_filename TEXT NOT NULL,
            source_sha256 TEXT NOT NULL,
            imported_by INTEGER,
            imported_at TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS paid_services_versions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            month TEXT NOT NULL,
            payload_json TEXT NOT NULL,
            source_filename TEXT NOT NULL,
            source_sha256 TEXT NOT NULL,
            imported_by INTEGER,
            imported_at TEXT NOT NULL,
            archived_by INTEGER,
            archived_at TEXT NOT NULL,
            decision TEXT NOT NULL
        );

        CREATE INDEX IF NOT EXISTS idx_paid_services_versions_month
        ON paid_services_versions(month, id);
        """
    )


def sha256(raw):
    return hashlib.sha256(raw).hexdigest()


def _clean_text(value):
    return str(value or "").replace("\xa0", " ").strip()


def _money(value):
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = _clean_text(value)
    if not text or text in {"-", "—"}:
        return 0.0
    text = text.replace(" ", "").replace(",", ".")
    try:
        return float(text)
    except ValueError as exc:
        raise ValueError(f"Не удалось распознать денежное значение «{value}».") from exc


def _date_value(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _clean_text(value)
    match = re.search(r"(\d{2})\.(\d{2})\.(\d{4})", text)
    if not match:
        return None
    day, month, year = map(int, match.groups())
    return date(year, month, day).isoformat()


def _fill_rgb(cell):
    fill = getattr(cell, "fill", None)
    if not fill or getattr(fill, "fill_type", None) != "solid":
        return ""
    color = getattr(fill, "fgColor", None)
    rgb = str(getattr(color, "rgb", "") or "").upper()
    return rgb


def _find_sheet(workbook):
    for ws in workbook.worksheets:
        header_values = {
            _clean_text(ws.cell(1, col).value)
            for col in range(1, min(ws.max_column, 12) + 1)
            if _clean_text(ws.cell(1, col).value)
        }
        if EXPECTED_HEADERS.issubset(header_values):
            return ws
    raise ValueError(
        "Не найден лист нового отчёта МИС с колонками "
        "«Сумма со скидкой / Оплачено / Задолженность»."
    )


def _provider_row(ws, row, section_fill):
    a = ws.cell(row, 1)
    name = _clean_text(a.value)
    if not name or name == "Услуги":
        return False
    if any(_clean_text(ws.cell(row, col).value) for col in (2, 3, 4)):
        return False
    fill = _fill_rgb(a)
    if not fill or fill in {section_fill, "FFFFFFFF", "00FFFFFF"}:
        return False
    values = [ws.cell(row, col).value for col in (5, 6, 7, 8)]
    return any(value not in (None, "", "-", "—") for value in values)


def _balance(opening, billed, paid, closing):
    return round(float(opening) + float(billed) - float(paid) - float(closing), 2)


def parse_file(raw, filename):
    if not raw:
        raise ValueError("Выбран пустой новый отчёт МИС.")
    if len(raw) > MAX_FILE_SIZE:
        raise ValueError("Размер файла превышает 25 МБ.")
    if not (raw.startswith(b"PK") or str(filename or "").lower().endswith(".xlsx")):
        raise ValueError("Новый отчёт МИС должен быть в формате XLSX.")

    try:
        workbook = load_workbook(BytesIO(raw), data_only=True, read_only=False)
    except Exception as exc:
        raise ValueError("Excel-файл повреждён или имеет неподдерживаемую структуру.") from exc

    try:
        ws = _find_sheet(workbook)
        services_row = None
        for row in range(1, ws.max_row + 1):
            if _clean_text(ws.cell(row, 1).value) == "Услуги":
                services_row = row
                break
        if services_row is None:
            raise ValueError("В новом отчёте МИС не найден раздел «Услуги».")

        section_fill = _fill_rgb(ws.cell(services_row, 1))
        totals = {
            "opening": round(_money(ws.cell(services_row, 5).value), 2),
            "billed": round(_money(ws.cell(services_row, 6).value), 2),
            "paid": round(_money(ws.cell(services_row, 7).value), 2),
            "closing": round(_money(ws.cell(services_row, 8).value), 2),
        }
        if abs(_balance(**totals)) > 0.02:
            raise ValueError(
                "Раздел «Услуги» не сходится: задолженность на начало + начислено "
                "− оплачено не равняется задолженности на конец."
            )

        providers = {}
        services = []
        report_dates = set()
        current_provider = None
        current_invoice = None
        current_invoice_date = None

        for row in range(services_row + 1, ws.max_row + 1):
            a = _clean_text(ws.cell(row, 1).value)
            b = ws.cell(row, 2).value
            c = _clean_text(ws.cell(row, 3).value)
            d = ws.cell(row, 4).value

            if _provider_row(ws, row, section_fill):
                current_provider = a
                current_invoice = None
                current_invoice_date = None
                item = {
                    "opening": round(_money(ws.cell(row, 5).value), 2),
                    "billed": round(_money(ws.cell(row, 6).value), 2),
                    "paid": round(_money(ws.cell(row, 7).value), 2),
                    "closing": round(_money(ws.cell(row, 8).value), 2),
                }
                if abs(_balance(**item)) > 0.02:
                    raise ValueError(f"Не сходится баланс по сотруднику «{current_provider}».")
                providers[current_provider] = item
                continue

            if not current_provider or not a:
                continue

            invoice_match = INVOICE_RE.match(a)
            if invoice_match:
                current_invoice = invoice_match.group(1)
                current_invoice_date = _date_value(invoice_match.group(2))
                if current_invoice_date:
                    report_dates.add(current_invoice_date)
                continue

            if a.lower().startswith("неоплаченные услуги на начало периода"):
                current_invoice = None
                current_invoice_date = None
                continue

            service_date = _date_value(b)
            if service_date and c:
                report_dates.add(service_date)
                try:
                    qty = float(d or 0)
                except (TypeError, ValueError):
                    qty = 0.0
                services.append(
                    {
                        "provider": current_provider,
                        "invoice": current_invoice,
                        "date": service_date,
                        "group": a,
                        "service": c,
                        "qty": qty,
                        "billed": round(_money(ws.cell(row, 6).value), 2),
                        "paid": round(_money(ws.cell(row, 7).value), 2),
                    }
                )

        if not providers:
            raise ValueError("В разделе «Услуги» не найдены строки сотрудников.")

        provider_totals = {
            key: round(sum(float(row[key] or 0) for row in providers.values()), 2)
            for key in ("opening", "billed", "paid", "closing")
        }
        for key in provider_totals:
            if abs(provider_totals[key] - totals[key]) > 0.02:
                raise ValueError(
                    "Сумма по сотрудникам не совпадает с итогом раздела «Услуги» "
                    f"для показателя {key}."
                )

        if not report_dates:
            raise ValueError("В новом отчёте МИС не найдены даты счетов/услуг.")
        months = {value[:7] for value in report_dates}
        if len(months) != 1:
            raise ValueError("Новый отчёт МИС должен содержать один календарный месяц.")
        month = next(iter(months))

        return {
            "version": 1,
            "month": month,
            "period": {"from": min(report_dates), "to": max(report_dates)},
            "totals": totals,
            "providers": providers,
            "services": services,
        }, ws.title
    finally:
        workbook.close()


def _provider_projection(payload, conn):
    maps = upload_reconcile.provider_maps(conn)
    known_nonclinical = set(ident_import.KNOWN_STAFF)
    direction_totals = {
        key: {"opening": 0.0, "billed": 0.0, "paid": 0.0, "closing": 0.0}
        for key in ("dent", "structure", "lab")
    }
    doctors = {"dent": {}, "structure": {}, "lab": {}}
    ignored = {}
    unknown = []

    for source_name, values in (payload.get("providers") or {}).items():
        direction = None
        display = source_name
        for candidate in ("dent", "structure", "lab"):
            if source_name in maps[candidate]:
                direction = candidate
                display = maps[candidate][source_name]
                break
        if direction is None and source_name in maps["ignore"]:
            direction = "ignore"
        if direction is None and source_name in known_nonclinical:
            direction = "ignore"

        if direction in {"dent", "structure", "lab"}:
            for key in direction_totals[direction]:
                direction_totals[direction][key] += float(values.get(key) or 0)
            doctors[direction][display] = round(
                doctors[direction].get(display, 0.0) + float(values.get("paid") or 0),
                2,
            )
        elif direction == "ignore":
            ignored[source_name] = dict(values)
        else:
            unknown.append(source_name)

    for direction in direction_totals.values():
        for key in direction:
            direction[key] = round(direction[key], 2)

    classified_paid = round(
        sum(direction_totals[key]["paid"] for key in ("dent", "structure", "lab")),
        2,
    )
    ignored_paid = round(sum(float(v.get("paid") or 0) for v in ignored.values()), 2)
    unknown_paid = round(
        sum(float((payload.get("providers") or {}).get(name, {}).get("paid") or 0) for name in unknown),
        2,
    )
    return {
        "directions": direction_totals,
        "doctors": doctors,
        "ignored": ignored,
        "unknown": sorted(unknown),
        "classified_paid": classified_paid,
        "ignored_paid": ignored_paid,
        "unknown_paid": unknown_paid,
    }


def _summary(payload, projection):
    totals = payload["totals"]
    directions = projection["directions"]
    return {
        "billed": totals["billed"],
        "paid": totals["paid"],
        "opening": totals["opening"],
        "closing": totals["closing"],
        "dent": directions["dent"]["paid"],
        "structure": directions["structure"]["paid"],
        "lab": directions["lab"]["paid"],
        "ignored": projection["ignored_paid"],
        "unclassified": projection["unknown_paid"],
    }


def _file_from_request():
    file_storage = request.files.get("paid_services")
    if not file_storage or not file_storage.filename:
        raise ValueError("Выберите новый отчёт МИС «Выручка по направлениям».")
    raw = file_storage.read()
    if not raw:
        raise ValueError("Выбран пустой новый отчёт МИС.")
    return file_storage.filename, raw


def _comparison(conn, payload):
    row = conn.execute(
        "SELECT payload_json,source_sha256 FROM paid_services_months WHERE month=?",
        (payload["month"],),
    ).fetchone()
    if not row:
        return "new"
    existing = json.loads(row["payload_json"])
    incoming = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    current = json.dumps(existing, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return "identical" if incoming == current else "conflict"


def latest_loaded_month(conn):
    init_schema(conn)
    row = conn.execute("SELECT MAX(month) FROM paid_services_months").fetchone()
    return str(row[0]) if row and row[0] else None


def _preview():
    perms = user_permissions(g.user)
    if "upload_services" not in perms:
        abort(403)

    filename, raw = _file_from_request()
    conn = db()
    init_schema(conn)
    upload_reconcile.init_schema(conn)
    upload_reconcile.refresh_runtime(conn, ident_import)

    payload, sheet = parse_file(raw, filename)
    projection = _provider_projection(payload, conn)
    state = _comparison(conn, payload)
    return {
        "status": "preview",
        "period": payload["period"],
        "month": payload["month"],
        "sheet": sheet,
        "counts": {
            "new": 1 if state == "new" else 0,
            "identical": 1 if state == "identical" else 0,
            "conflict": 1 if state == "conflict" else 0,
            "historical": 0,
            "removed": 0,
        },
        "requires_choice": state == "conflict",
        "requires_provider_mapping": bool(projection["unknown"]),
        "unknown_providers": projection["unknown"],
        "can_replace": "upload_replace" in perms,
        "summary": _summary(payload, projection),
    }


def _provider_decisions_from_request():
    raw = request.form.get("provider_decisions", "").strip()
    if not raw:
        return {}
    try:
        value = json.loads(raw)
    except ValueError as exc:
        raise ValueError("Не удалось прочитать классификацию новых сотрудников.") from exc
    return value if isinstance(value, dict) else {}


def _archive_existing(conn, row, actor_id, decision):
    conn.execute(
        """
        INSERT INTO paid_services_versions(
            month,payload_json,source_filename,source_sha256,
            imported_by,imported_at,archived_by,archived_at,decision
        ) VALUES(?,?,?,?,?,?,?,?,?)
        """,
        (
            row["month"],
            row["payload_json"],
            row["source_filename"],
            row["source_sha256"],
            row["imported_by"],
            row["imported_at"],
            actor_id,
            iso_now(),
            decision,
        ),
    )


def _commit(decision=None):
    perms = user_permissions(g.user)
    if "upload_services" not in perms:
        abort(403)

    filename, raw = _file_from_request()
    conn = db()
    init_schema(conn)
    upload_reconcile.init_schema(conn)
    upload_reconcile.refresh_runtime(conn, ident_import)

    payload, sheet = parse_file(raw, filename)
    decisions = _provider_decisions_from_request()
    projection = _provider_projection(payload, conn)

    if projection["unknown"]:
        if not decisions:
            raise ValueError("Сначала классифицируйте новых сотрудников из нового отчёта МИС.")
        if "upload_replace" not in perms:
            abort(403)
        upload_reconcile.resolve_providers(conn, decisions, g.user["id"])
        upload_reconcile.refresh_runtime(conn, ident_import)
        projection = _provider_projection(payload, conn)
        if projection["unknown"]:
            raise ValueError("Не для всех новых сотрудников выбрано направление.")

    state = _comparison(conn, payload)
    if state == "conflict" and decision not in {"keep_old", "use_new"}:
        raise ValueError(
            "За этот месяц уже загружена другая версия нового отчёта МИС. "
            "Выберите «Сохранить старые» или «Загрузить новые»."
        )
    if decision == "use_new" and "upload_replace" not in perms:
        abort(403)

    if state == "identical":
        return {
            "status": "duplicate",
            "message": "Новый отчёт МИС совпадает с сохранённым; ничего не изменено.",
            "summary": _summary(payload, projection),
        }
    if state == "conflict" and decision == "keep_old":
        return {
            "status": "duplicate",
            "message": "Сохранена ранее загруженная версия нового отчёта МИС.",
            "summary": _summary(payload, projection),
        }

    backup = None
    now = iso_now()
    source_sha = sha256(raw)
    existing = conn.execute(
        "SELECT * FROM paid_services_months WHERE month=?",
        (payload["month"],),
    ).fetchone()
    if existing and decision == "use_new":
        backup = upload_reconcile.create_db_backup("paid-services-reimport")

    conn.execute("BEGIN")
    try:
        if existing:
            _archive_existing(conn, existing, g.user["id"], decision or "use_new")
            conn.execute("DELETE FROM paid_services_months WHERE month=?", (payload["month"],))
        conn.execute(
            """
            INSERT INTO paid_services_months(
                month,payload_json,source_filename,source_sha256,imported_by,imported_at
            ) VALUES(?,?,?,?,?,?)
            """,
            (
                payload["month"],
                json.dumps(payload, ensure_ascii=False, separators=(",", ":")),
                filename,
                source_sha,
                g.user["id"],
                now,
            ),
        )
        conn.commit()
    except Exception:
        conn.rollback()
        raise

    audit(
        "paid_services_reconciled",
        target_user_id=g.user["id"],
        details=(
            f"file={filename}; sheet={sheet}; month={payload['month']}; "
            f"decision={decision or 'append'}; paid={payload['totals']['paid']}; "
            f"providers={len(payload['providers'])}; backup={backup or ''}; "
            "reports_updated=0"
        ),
    )

    return {
        "status": "replaced" if existing else "imported",
        "message": (
            "Новая версия отчёта МИС сохранена отдельно для проверки. "
            "Действующие управленческие отчёты этим действием не пересчитываются."
        ),
        "backup": backup,
        "summary": _summary(payload, projection),
    }


def register_paid_services_upload(app):
    @app.post("/api/uploads/paid-services/preview")
    @permission_required("section5")
    def paid_services_preview():
        require_csrf()
        try:
            return jsonify(_preview())
        except ValueError as exc:
            return jsonify({"status": "error", "message": str(exc)}), 400

    @app.post("/api/uploads/paid-services/commit")
    @permission_required("section5")
    def paid_services_commit():
        require_csrf()
        try:
            return jsonify(_commit(request.form.get("decision") or None))
        except ValueError as exc:
            return jsonify({"status": "error", "message": str(exc)}), 409
