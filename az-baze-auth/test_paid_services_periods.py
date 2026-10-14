import io
import json
import os
import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from flask import Flask, g
from openpyxl import Workbook, load_workbook

import paid_services

with patch.dict(os.environ, {"AZBAZE_SECRET_KEY": "paid-period-tests-only"}):
    import paid_services_upload as uploads


def report_bytes(dates=("05.01.2026", "27.06.2026"), paid=None, provider="Dent D.", period=None):
    """Synthetic source: provider paid includes 60 paid against opening debt."""
    billed = 100 * len(dates)
    paid = billed + 60 if paid is None else paid
    closing = 100 + billed - paid
    if period is None and dates:
        parsed_dates = [__import__("datetime").datetime.strptime(value, "%d.%m.%Y") for value in dates]
        first, last = min(parsed_dates), max(parsed_dates)
        if first.year == last.year and first.month == last.month:
            period = (first.replace(day=1), last)
        else:
            import calendar
            period = (first.replace(day=1), last.replace(day=calendar.monthrange(last.year, last.month)[1]))
    wb = Workbook()
    ws = wb.active
    ws.title = "Выручка по направлениям"
    period_label = ""
    if period:
        start, end = period
        if not hasattr(start, "strftime"):
            start = __import__("datetime").datetime.strptime(str(start), "%d.%m.%Y")
        if not hasattr(end, "strftime"):
            end = __import__("datetime").datetime.strptime(str(end), "%d.%m.%Y")
        period_label = f"Период {start.strftime('%d.%m.%Y')} — {end.strftime('%d.%m.%Y')}"
    ws.append([
        "Группа услуг", "Услуги", None, None,
        "Задолженность на нач. периода", "Сумма со скидкой",
        "Оплачено", "Задолженность на конец периода", period_label,
    ])
    ws.append(["Услуги", None, None, None, 100, billed, paid, closing])
    ws.append([provider, None, None, None, 100, billed, paid, closing])
    ws.append(["Неоплаченные услуги на начало периода", None, None, None, 100, 0, 60, 40])
    for number, day in enumerate(dates, 1):
        ws.append(["Синтетический пациент", None, None, None, 0, 100, 100, 0])
        ws.append([f"Счет №{number} от {day} 12:00:00", None, None, None, 0, 100, 100, 0])
        ws.append(["Терапия", day + " 12:00", "Тестовая услуга", 1, 0, 100, 100, 0])
    stream = io.BytesIO()
    wb.save(stream)
    wb.close()
    return stream.getvalue()


def parse_period(raw=None, start="2026-01-01", end="2026-06-30"):
    return paid_services.parse_bytes(
        report_bytes() if raw is None else raw,
        "period.xlsx", period_start=start, period_end=end,
    )


class PaidPeriodParserTests(unittest.TestCase):
    def test_nonfinite_money_cannot_bypass_source_balance_checks(self):
        for value in ("NaN", "Infinity", "-Infinity"):
            with self.subTest(value=value):
                wb = load_workbook(io.BytesIO(report_bytes()))
                wb.active["G2"] = value
                stream = io.BytesIO()
                wb.save(stream)
                wb.close()
                with self.assertRaisesRegex(ValueError, "конечным числом"):
                    parse_period(stream.getvalue())
        for value in (float("nan"), float("inf"), float("-inf")):
            with self.subTest(value=str(value)), self.assertRaisesRegex(ValueError, "конечным числом"):
                paid_services._num(value)

    def test_half_year_keeps_source_paid_without_month_allocation(self):
        parsed = parse_period()
        self.assertEqual((parsed["period_start"], parsed["period_end"]), ("2026-01-01", "2026-06-30"))
        self.assertEqual(parsed["month"], "")
        self.assertEqual(parsed["providers"]["Dent D."]["paid"], 260)
        self.assertEqual(sum(row["paid_amount"] for row in parsed["items"]), 200)
        compact = paid_services.compact_report(parsed)
        self.assertEqual(compact["version"], 3)
        self.assertNotIn("month", compact)
        self.assertNotIn("as_of_date", compact)
        self.assertEqual(compact["providers"]["Dent D."]["paid"], 260)
        self.assertNotIn("Синтетический пациент", json.dumps(compact, ensure_ascii=False))
        self.assertEqual(len(compact["visit_links"]), 2)

    def test_cross_year_and_partial_month_use_interval_storage(self):
        cases = [
            (("29.12.2025", "03.01.2026"), "2025-12-20", "2026-01-10"),
            (("16.06.2026", "27.06.2026"), "2026-06-15", "2026-06-30"),
        ]
        for dates, start, end in cases:
            with self.subTest(start=start, end=end):
                parsed = parse_period(report_bytes(dates), start, end)
                self.assertFalse(paid_services.is_monthly_report(parsed))
                self.assertEqual(paid_services.compact_report(parsed)["version"], 3)

    def test_same_month_without_period_label_keeps_actual_source_dates(self):
        raw = report_bytes(("05.07.2026", "31.07.2026"), period=False)
        inferred = paid_services.compact_report(paid_services.parse_bytes(raw, "july-slice.xlsx"))
        explicit = paid_services.compact_report(parse_period(raw, "2026-07-01", "2026-07-31"))
        self.assertEqual(inferred["version"], 3)
        self.assertEqual(inferred["period_start"], "2026-07-05")
        self.assertEqual(inferred["period_end"], "2026-07-31")
        self.assertNotIn("month", inferred)
        self.assertEqual(explicit["version"], 2)
        self.assertEqual(explicit["month"], "2026-07")
        self.assertNotIn("period_start", explicit)

    def test_explicit_end_can_include_days_without_services(self):
        parsed = parse_period(report_bytes(("05.06.2026", "27.06.2026")), "2026-06-01", "2026-06-30")
        self.assertEqual(parsed["period_end"], "2026-06-30")
        self.assertEqual(paid_services.compact_report(parsed)["as_of_date"], "2026-06-30")

    def test_period_is_read_from_workbook_and_debt_only_needs_a_period_source(self):
        parsed = paid_services.parse_bytes(report_bytes(), "half-year.xlsx")
        self.assertEqual((parsed["period_start"], parsed["period_end"]), ("2026-01-01", "2026-06-30"))
        with self.assertRaisesRegex(ValueError, "не найдены даты периода"):
            paid_services.parse_bytes(report_bytes(()))
        parsed = parse_period(report_bytes(()))
        self.assertEqual(parsed["items"], [])
        self.assertEqual(parsed["providers"]["Dent D."]["paid"], 60)

    def test_multi_month_without_period_label_uses_dates_inside_file(self):
        raw = report_bytes(period=False)
        parsed = paid_services.parse_bytes(raw, "dated-rows-only.xlsx")
        self.assertEqual((parsed["period_start"], parsed["period_end"]), ("2026-01-05", "2026-06-27"))

    def test_invalid_incomplete_reversed_or_incompatible_dates_are_rejected(self):
        cases = [
            ("2026-01-01", None), (None, "2026-06-30"),
            ("2026-02-30", "2026-06-30"), ("2026-6-1", "2026-06-30"),
            ("2026-07-01", "2026-06-30"),
            ("2026-01-06", "2026-06-30"), ("2026-01-01", "2026-06-26"),
        ]
        for start, end in cases:
            with self.subTest(start=start, end=end), self.assertRaises(ValueError):
                parse_period(start=start, end=end)


class PaidPeriodStorageTests(unittest.TestCase):
    def setUp(self):
        self.conn = sqlite3.connect(":memory:")
        self.conn.row_factory = sqlite3.Row
        paid_services.init_schema(self.conn)
        self.addCleanup(self.conn.close)

    def store(self, parsed, sha="first", decision="append"):
        return paid_services.store_snapshot(self.conn, parsed, sha + ".xlsx", sha, 7, "now", decision)

    def test_overlap_and_same_end_are_distinct_sources_and_leave_months_unchanged(self):
        june = parse_period(report_bytes(("05.06.2026", "27.06.2026")), "2026-06-01", "2026-06-30")
        self.assertEqual(self.store(june), "imported")
        self.conn.commit()
        original = tuple(self.conn.execute("SELECT * FROM service_payment_snapshots").fetchone())
        self.assertEqual(self.store(parse_period(), "half-year"), "imported")
        overlap = parse_period(report_bytes(("05.04.2026", "27.09.2026")), "2026-04-01", "2026-09-30")
        self.assertEqual(self.store(overlap, "overlap"), "imported")
        self.assertEqual(tuple(self.conn.execute("SELECT * FROM service_payment_snapshots").fetchone()), original)
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM service_payment_periods").fetchone()[0], 2)
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM service_payment_versions").fetchone()[0], 0)
        self.assertEqual(paid_services.latest_loaded_date(self.conn), "2026-09-30")
        self.assertEqual(paid_services.snapshot_months(self.conn), ["2026-06"])

    def test_identical_period_is_noop_and_replacement_archives_exact_source(self):
        original = parse_period()
        self.assertEqual(self.store(original), "imported")
        self.assertEqual(self.store(original, "different-file-hash"), "identical")
        changed = parse_period(report_bytes(paid=270))
        with self.assertRaisesRegex(ValueError, "другая версия"):
            self.store(changed, "second")
        self.assertEqual(self.store(changed, "second", "use_new"), "replaced")
        archive = self.conn.execute("SELECT * FROM service_payment_period_versions").fetchone()
        self.assertEqual(archive["source_sha256"], "first")
        self.assertEqual(json.loads(archive["payload_json"])["providers"]["Dent D."]["paid"], 260)
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM service_payment_period_versions").fetchone()[0], 1)

    def test_caller_rollback_reverts_snapshot_and_archive_for_both_storage_types(self):
        cases = [
            (parse_period(), parse_period(report_bytes(paid=270)), "service_payment_periods", "service_payment_period_versions"),
            (parse_period(report_bytes(("05.06.2026",)), "2026-06-01", "2026-06-30"),
             parse_period(report_bytes(("05.06.2026",), paid=170), "2026-06-01", "2026-06-30"),
             "service_payment_snapshots", "service_payment_versions"),
        ]
        for first, changed, table, archive in cases:
            with self.subTest(table=table):
                self.store(first)
                self.conn.commit()
                before = [tuple(row) for row in self.conn.execute("SELECT * FROM " + table)]
                self.conn.execute("BEGIN IMMEDIATE")
                self.store(changed, "second", "use_new")
                self.assertTrue(self.conn.in_transaction)
                self.conn.rollback()
                self.assertEqual([tuple(row) for row in self.conn.execute("SELECT * FROM " + table)], before)
                self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM " + archive).fetchone()[0], 0)


class PaidPeriodUploadTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix="paid-period-api-")
        self.addCleanup(temporary.cleanup)
        self.path = Path(temporary.name) / "auth.db"
        self.conn = sqlite3.connect(self.path)
        self.conn.row_factory = sqlite3.Row
        self.addCleanup(self.conn.close)
        paid_services.init_schema(self.conn)
        uploads.upload_reconcile.init_schema(self.conn)
        self.conn.executescript("""
            CREATE TABLE report_data(date TEXT PRIMARY KEY,payload TEXT);
            INSERT INTO report_data VALUES('2026-09-30','{"cashTotal":70,"cashOOO":30,"cashIP":40}');
            CREATE TABLE report_blobs(key TEXT PRIMARY KEY,payload TEXT);
            INSERT INTO report_blobs VALUES('analytics','{"approved":true}');
            CREATE TABLE daily_uploads(data_date TEXT PRIMARY KEY,normalized_json TEXT);
            INSERT INTO daily_uploads VALUES('2026-09-30','{"primary":12}');
        """)
        self.protected = self.protected_state()
        self.perms = {"upload_services", "upload_replace"}
        for patcher in (
            patch.object(uploads, "db", return_value=self.conn),
            patch.object(uploads, "audit"),
            patch.object(uploads, "iso_now", return_value="2026-10-12T09:00:00+00:00"),
            patch.object(uploads, "user_permissions", side_effect=lambda user: self.perms),
            patch.object(uploads, "permission_required", side_effect=lambda permission: lambda function: function),
            patch.object(uploads, "require_csrf"),
            patch.object(uploads.ident_import, "DENTISTS", {"Dent D.": "Dent Doctor"}),
            patch.object(uploads.ident_import, "STRUCTURE_DOCTORS", {}),
            patch.object(uploads.ident_import, "LAB_DOCTORS", {}),
            patch.object(uploads.ident_import, "KNOWN_STAFF", {"Dent D."}),
            patch.object(uploads.cash_payments, "_PROVIDER_MAPS", {"dent": {}, "structure": {}, "lab": {}}),
            patch.object(uploads.upload_reconcile, "DB_PATH", self.path),
        ):
            patcher.start()
            self.addCleanup(patcher.stop)
        self.app = Flask(__name__)
        self.app.config["TESTING"] = True

        @self.app.before_request
        def set_user():
            g.user = {"id": 7}

        uploads.register_paid_services_upload(self.app)
        self.client = self.app.test_client()

    def protected_state(self):
        return {
            table: [tuple(row) for row in self.conn.execute("SELECT * FROM " + table)]
            for table in ("report_data", "report_blobs", "daily_uploads")
        }

    def post(self, action, raw=None, **fields):
        data = {
            **fields,
            "paid_services": (io.BytesIO(report_bytes() if raw is None else raw), "period.xlsx"),
        }
        return self.client.post("/api/uploads/paid-services/" + action, data=data)

    def import_first(self):
        preview = self.post("preview").get_json()
        self.assertEqual(preview["revision"], "new")
        result = self.post("commit", decision="append", expected_revision=preview["revision"])
        self.assertEqual(result.status_code, 200, result.get_json())
        self.assertEqual(result.get_json()["status"], "imported")
        return preview

    def test_preview_of_unknown_provider_is_read_only_and_reports_exact_period(self):
        before = list(self.conn.iterdump())
        result = self.post("preview", report_bytes(provider="Unknown U."))
        self.assertEqual(result.status_code, 200)
        preview = result.get_json()
        self.assertEqual(preview["period"], {"from": "2026-01-01", "to": "2026-06-30"})
        self.assertEqual(preview["unknown_providers"], ["Unknown U."])
        self.assertEqual(list(self.conn.iterdump()), before)

    def test_replacement_uses_preview_revision_backup_and_archive(self):
        self.import_first()
        changed = report_bytes(paid=270)
        preview = self.post("preview", changed).get_json()
        self.assertTrue(preview["requires_choice"])
        self.assertTrue(preview["requires_review"])
        self.assertTrue(preview["changes_detected"])
        self.assertEqual(preview["overlaps"][0]["relation"], "exact")
        result = self.post("commit", changed, decision="use_new", expected_revision=preview["revision"])
        self.assertEqual(result.status_code, 200, result.get_json())
        payload = result.get_json()
        self.assertEqual(payload["status"], "replaced")
        self.assertEqual(payload["summary"]["dent"], 270)
        self.assertTrue(Path(payload["backup"]).is_file())
        with sqlite3.connect(payload["backup"]) as backup:
            old = json.loads(backup.execute("SELECT payload_json FROM service_payment_periods").fetchone()[0])
            self.assertEqual(old["providers"]["Dent D."]["paid"], 260)
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM service_payment_period_versions").fetchone()[0], 1)
        self.assertEqual(self.protected_state(), self.protected)

    def test_partial_overlap_is_reported_before_append(self):
        self.import_first()
        raw = report_bytes(("05.04.2026", "27.09.2026"))
        preview = self.post("preview", raw).get_json()
        self.assertFalse(preview["requires_choice"])
        self.assertTrue(preview["requires_review"])
        self.assertEqual(preview["counts"]["overlap"], 1)
        self.assertEqual(preview["overlaps"][0]["relation"], "partial")
        self.assertEqual(preview["period"], {"from": "2026-04-01", "to": "2026-09-30"})

    def test_stale_preview_cannot_overwrite_a_later_import(self):
        self.import_first()
        desired = report_bytes(paid=270)
        preview = self.post("preview", desired).get_json()
        paid_services.store_snapshot(self.conn, parse_period(report_bytes(paid=280)), "other.xlsx", "other", 8, "later", "use_new")
        self.conn.commit()
        before = list(self.conn.iterdump())
        result = self.post("commit", desired, decision="use_new", expected_revision=preview["revision"])
        self.assertEqual(result.status_code, 409)
        self.assertIn("уже изменился", result.get_json()["message"])
        self.assertEqual(list(self.conn.iterdump()), before)

    def test_keep_old_does_not_classify_unknown_names_from_rejected_file(self):
        self.import_first()
        raw = report_bytes(provider="Unknown U.")
        preview = self.post("preview", raw).get_json()
        before = list(self.conn.iterdump())
        result = self.post("commit", raw, decision="keep_old", expected_revision=preview["revision"])
        self.assertEqual(result.status_code, 200)
        self.assertEqual(result.get_json()["status"], "duplicate")
        self.assertEqual(list(self.conn.iterdump()), before)

    def test_classification_and_snapshot_roll_back_together(self):
        raw = report_bytes(provider="Unknown U.")
        decisions = json.dumps({"Unknown U.": {"direction": "dent", "display_name": "New Doctor"}})
        before = list(self.conn.iterdump())
        with patch.object(paid_services, "store_snapshot", side_effect=RuntimeError("storage failure")):
            with self.assertRaisesRegex(RuntimeError, "storage failure"):
                self.post("commit", raw, expected_revision="new", provider_decisions=decisions)
        self.assertEqual(list(self.conn.iterdump()), before)
        self.assertFalse(self.conn.in_transaction)

    def test_classification_cannot_rewrite_other_registry_entries(self):
        self.conn.execute(
            "INSERT INTO provider_registry VALUES(?,?,?,NULL,?,?,?,?)",
            ("Dent D.", "Dent Doctor", "dent", "before", "before", 1, "before"),
        )
        self.conn.commit()
        before = tuple(self.conn.execute("SELECT * FROM provider_registry WHERE source_name='Dent D.'").fetchone())
        decisions = json.dumps({
            "Unknown U.": {"direction": "structure", "display_name": "New Doctor"},
            "Dent D.": {"direction": "ignore", "display_name": "Changed"},
        })
        result = self.post("commit", report_bytes(provider="Unknown U."), expected_revision="new", provider_decisions=decisions)
        self.assertEqual(result.status_code, 200, result.get_json())
        self.assertEqual(result.get_json()["summary"]["structure"], 260)
        self.assertEqual(tuple(self.conn.execute("SELECT * FROM provider_registry WHERE source_name='Dent D.'").fetchone()), before)
        self.assertEqual(self.protected_state(), self.protected)

    def test_replacement_and_classification_require_existing_permission(self):
        self.import_first()
        raw = report_bytes(paid=270)
        revision = self.post("preview", raw).get_json()["revision"]
        self.perms = {"upload_services"}
        before = list(self.conn.iterdump())
        self.assertEqual(self.post("commit", raw, decision="use_new", expected_revision=revision).status_code, 403)
        self.assertEqual(self.post("commit", report_bytes(("05.07.2026", "31.07.2026"), provider="Unknown U."),
                                   expected_revision="new",
                                   provider_decisions=json.dumps({"Unknown U.": "dent"})).status_code, 403)
        self.assertEqual(list(self.conn.iterdump()), before)

    def test_invalid_classification_is_rejected_without_partial_writes(self):
        before = list(self.conn.iterdump())
        result = self.post("commit", report_bytes(provider="Unknown U."), expected_revision="new",
                           provider_decisions=json.dumps({"Unknown U.": "not-a-direction"}))
        self.assertEqual(result.status_code, 409)
        self.assertEqual(list(self.conn.iterdump()), before)


if __name__ == "__main__":
    unittest.main()
