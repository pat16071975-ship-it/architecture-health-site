import io
import json
import os
import sqlite3
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

os.environ.setdefault("AZBAZE_SECRET_KEY", "test-secret")
os.environ["AZBAZE_OWNER_EMAIL"] = "owner@example.com"

import app as app_module
import nn_normalize
import nn_upload
import nn_reports


class NNReportsTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        root = Path(self.tmp.name)
        app_module.DB_PATH = root / "auth.db"
        app_module.SITE_ROOT = root / "site"
        nn_upload.DB_PATH = app_module.DB_PATH
        nn_upload.UPLOAD_ROOT = root / "nn-uploads"
        app_module.SITE_ROOT.mkdir(parents=True, exist_ok=True)
        (app_module.SITE_ROOT / "index.html").write_text(
            "<!doctype html><html><body><nav></nav></body></html>",
            encoding="utf-8",
        )
        app_module.init_db_file(app_module.DB_PATH)

        conn = sqlite3.connect(app_module.DB_PATH)
        conn.row_factory = sqlite3.Row
        conn.execute("PRAGMA foreign_keys=ON")
        conn.executescript(
            """
            CREATE TABLE clinics (
                id INTEGER PRIMARY KEY,
                name TEXT NOT NULL,
                address TEXT NOT NULL DEFAULT '',
                status TEXT NOT NULL
            );
            """
        )
        now = app_module.iso_now()
        conn.executemany(
            """INSERT INTO users(
                id,email,full_name,password_hash,is_admin,active,must_change_password,
                failed_attempts,locked_until,created_at,updated_at
            ) VALUES(?,?,?,?,?,1,0,0,NULL,?,?)""",
            [
                (1, "owner@example.com", "Owner", "x", 1, now, now),
                (2, "reports@example.com", "Reports", "x", 0, now, now),
                (3, "upload@example.com", "Upload", "x", 0, now, now),
            ],
        )
        conn.executemany(
            "INSERT INTO permissions(user_id,section) VALUES(?,?)",
            [(2, "nn_reports"), (3, "nn_upload")],
        )
        conn.execute(
            "INSERT INTO clinics(id,name,address,status) VALUES(11,'Архитектура здоровья','AZ address','active')"
        )
        conn.commit()
        conn.close()

        nn_upload._init_schema()
        conn = sqlite3.connect(app_module.DB_PATH)
        conn.row_factory = sqlite3.Row
        conn.execute(
            "INSERT INTO nn_upload_batches(id,clinic_id,bundle_sha256,status,uploaded_by,uploaded_at) VALUES(101,1,'hash','ready',1,?)",
            (now,),
        )
        conn.commit()
        nn_normalize._init_schema(conn)

        payload = {
            "version": 1,
            "period": {"start": "2026-06-01", "end": "2026-07-31"},
            "source_control": {},
            "patients": [
                {"key": "card:1", "name": "Пациент Один", "dob": "1980-01-01", "iin": "", "chart": "1", "phone": "7000000001", "note": ""},
                {"key": "card:2", "name": "Пациент Два", "dob": "1981-01-01", "iin": "", "chart": "2", "phone": "7000000002", "note": ""},
            ],
            "visits": [
                {"patient_key": "card:1", "date": "2026-06-01", "start": "09:00", "end": "09:30", "doctor": "Врач А", "patient": "Пациент Один", "chart": "1", "phone": "7000000001", "iin": "", "repeat_field": "", "note": "", "visit_type": "Амбулаторно", "help_type": "", "services_text": "Прием хирурга флеболога с УЗИ", "amount": 4900},
                {"patient_key": "card:1", "date": "2026-06-10", "start": "09:00", "end": "09:30", "doctor": "Врач А", "patient": "Пациент Один", "chart": "1", "phone": "7000000001", "iin": "", "repeat_field": "", "note": "повторный контроль", "visit_type": "Амбулаторно", "help_type": "", "services_text": "", "amount": 0},
                {"patient_key": "card:1", "date": "2026-06-20", "start": "09:00", "end": "09:30", "doctor": "Врач А", "patient": "Пациент Один", "chart": "1", "phone": "7000000001", "iin": "", "repeat_field": "", "note": "контроль после операции", "visit_type": "Амбулаторно", "help_type": "", "services_text": "", "amount": 0},
                {"patient_key": "card:2", "date": "2026-06-02", "start": "10:00", "end": "10:30", "doctor": "Врач Б", "patient": "Пациент Два", "chart": "2", "phone": "7000000002", "iin": "", "repeat_field": "", "note": "", "visit_type": "Амбулаторно", "help_type": "", "services_text": "Прием хирурга флеболога с УЗИ", "amount": 0},
                {"patient_key": "card:1", "date": "2026-06-05", "start": "08:00", "end": "09:00", "doctor": "Врач А", "patient": "Пациент Один", "chart": "1", "phone": "7000000001", "iin": "", "repeat_field": "", "note": "ЭВЛК + МФ", "visit_type": "Амбулаторно", "help_type": "", "services_text": "Лазерное лечение варикоза категория B", "amount": 300000},
                {"patient_key": "card:2", "date": "2026-07-02", "start": "10:00", "end": "10:30", "doctor": "Врач Б", "patient": "Пациент Два", "chart": "2", "phone": "7000000002", "iin": "", "repeat_field": "повторный", "note": "", "visit_type": "Амбулаторно", "help_type": "", "services_text": "", "amount": 0},
            ],
            "medical_records": [
                {"date": "2026-06-05", "doctor": "Врач А", "patient": "Пациент Один", "dob": "1980-01-01", "chart": "1", "complaint": "", "objective": "", "direction": "", "assignment": "", "recommendation": "", "treatment": "Протокол манипуляции: Эндоваскулярная лазерная коагуляция + минифлебэктомия"},
                {"date": "2026-06-15", "doctor": "Врач Б", "patient": "Пациент Два", "dob": "1981-01-01", "chart": "2", "complaint": "", "objective": "", "direction": "", "assignment": "", "recommendation": "", "treatment": "Протокол манипуляции: Пенная Склеротерапия"},
            ],
            "service_control_rows": [],
            "deleted_appointments": [],
            "primaries": [],
            "repeats": [],
            "treatments": [],
            "suspicious": [],
            "monthly": [],
            "doctor_metrics": [],
            "key_metrics": {},
        }
        conn.execute(
            "INSERT INTO nn_normalized_batches(batch_id,clinic_id,version,period_start,period_end,payload_json,normalized_at) VALUES(101,1,1,'2026-06-01','2026-07-31',?,?)",
            (json.dumps(payload, ensure_ascii=False), now),
        )
        conn.commit()
        conn.close()

        self.web = app_module.create_app()
        self.web.config.update(TESTING=True)

    def tearDown(self):
        self.tmp.cleanup()

    def client_for(self, user_id):
        client = self.web.test_client()
        with client.session_transaction() as session:
            session["user_id"] = user_id
            session["csrf"] = "test-csrf"
        return client

    def test_reports_page_has_all_eight_reports_and_private_brand(self):
        response = self.client_for(2).get("/nn/reports/")
        html = response.get_data(as_text=True)
        self.assertEqual(response.status_code, 200)
        for label in (
            "Ключевые показатели",
            "Помесячная сводка",
            "Врачи — сравнение",
            "Врачи",
            "Подозрительные",
            "Первичные",
            "Лечение",
            "Первичные / повторные пациенты",
        ):
            self.assertIn(label, html)
        self.assertIn("Варикоза нет - KZ", html)
        self.assertNotIn("Архитектура здоровья", html)
        self.assertIn("Клиника 1 (Толи Бе)", html)
        self.assertIn("Клиника 2 (другая)", html)
        self.assertIn("Не оплатил первичный, но продолжил ходить дальше", html)
        self.assertIn(".nn-shell{width:min(1420px,97vw)}", html)
        self.assertIn(".nn-table.compare{min-width:1160px}", html)
        self.assertIn("Оплативших<br>первичный", html)
        self.assertIn(".nn-metric span{font-size:16px", html)
        self.assertIn(".nn-table th{position:sticky;top:0;background:#f8f3eb;z-index:3;white-space:normal;font-size:11px", html)
        self.assertIn("Не пришли повторно после неоплаченного первичного", html)
        self.assertIn("Скачать подробный отчёт Excel", html)

    def test_latest_month_and_available_year_month_context(self):
        response = self.client_for(2).get("/api/nn/reports/key_metrics?clinic_id=1")
        data = response.get_json()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(data["period"]["start"], "2026-07-01")
        self.assertEqual(data["context"]["years"], [2026])
        self.assertEqual(data["context"]["months_by_year"]["2026"], [6, 7])

    def test_key_metrics_does_not_duplicate_unique_repeat_people_report(self):
        response = self.client_for(2).get(
            "/api/nn/reports/key_metrics?clinic_id=1&year=2026&month=6"
        )
        data = response.get_json()
        labels = [item[0] for item in data["items"]]
        self.assertNotIn("Уникальных повторных пациентов", labels)

    def test_suspicious_is_grouped_month_then_date(self):
        response = self.client_for(2).get(
            "/api/nn/reports/suspicious?clinic_id=1&year=2026&month=6"
        )
        data = response.get_json()
        self.assertEqual(response.status_code, 200)
        self.assertTrue(data["groups"])
        self.assertEqual(data["groups"][0]["month"], "2026-06")
        self.assertTrue(data["groups"][0]["days"])
        self.assertEqual(data["groups"][0]["days"][0]["date"], "2026-06-15")
        self.assertTrue(data["groups"][0]["days"][0]["rows"])
        self.assertTrue(data["levels"])
        self.assertEqual(data["levels"][0]["key"], "high")
        self.assertEqual(data["levels"][0]["label"], "Высокая вероятность несоответствия")

    def test_suspicious_confidence_levels_follow_approved_logic(self):
        self.assertEqual(
            nn_reports._suspicious_level({"treatment": "protocol", "repeat": "2026-06-10"}),
            ("high", "Высокая вероятность несоответствия"),
        )
        self.assertEqual(
            nn_reports._suspicious_level({"treatment": "protocol", "repeat": ""}),
            ("review", "Требует проверки"),
        )
        self.assertEqual(
            nn_reports._suspicious_level({"treatment": "", "repeat": "2026-06-10"}),
            ("insufficient", "Недостаточно данных"),
        )

    def test_primary_repeat_report_finds_unpaid_primary_who_returned_later(self):
        response = self.client_for(2).get(
            "/api/nn/reports/primary_repeat?clinic_id=1&date_from=2026-06-01&date_to=2026-06-30"
        )
        data = response.get_json()
        self.assertEqual(response.status_code, 200)
        # Patient 2 has an unpaid primary on 02.06 and a repeat on 02.07.
        # Patient 1 paid the primary and must not be part of this report.
        self.assertEqual(data["unpaid_primary_people"], 1)
        self.assertEqual(data["continued_people"], 1)
        self.assertEqual(data["no_repeat_people"], 0)
        self.assertEqual(data["continued_share"], 100.0)
        self.assertEqual(len(data["rows"]), 1)
        self.assertEqual(data["rows"][0]["patient_name"], "Пациент Два")
        self.assertEqual(data["rows"][0]["chart"], "2")
        self.assertEqual(data["rows"][0]["primary_date"], "2026-06-02")
        self.assertEqual(data["rows"][0]["primary_doctor"], "Врач Б")
        self.assertEqual(data["rows"][0]["first_repeat_date"], "2026-07-02")
        self.assertEqual(data["rows"][0]["first_repeat_doctor"], "Врач Б")
        self.assertEqual(data["rows"][0]["repeat_count"], 1)

    def test_doctor_filter_returns_one_doctor(self):
        response = self.client_for(2).get(
            "/api/nn/reports/doctors?clinic_id=1&year=2026&month=6&doctor=Врач%20А"
        )
        data = response.get_json()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(data["doctors"]), 1)
        self.assertEqual(data["doctors"][0]["doctor"], "Врач А")

    def test_unmapped_clinic_is_rejected(self):
        response = self.client_for(2).get("/api/nn/reports/key_metrics?clinic_id=999")
        self.assertEqual(response.status_code, 403)

    def test_upload_only_user_cannot_access_reports(self):
        self.assertEqual(self.client_for(3).get("/nn/reports/").status_code, 403)
        self.assertEqual(self.client_for(3).get("/api/nn/reports/key_metrics?clinic_id=1").status_code, 403)

    def test_primary_export_is_xlsx_and_respects_paid_tab(self):
        response = self.client_for(2).get(
            "/nn/reports/export/primaries.xlsx?clinic_id=1&year=2026&month=6&tab=paid"
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response.mimetype)
        wb = load_workbook(io.BytesIO(response.data), read_only=True, data_only=True)
        try:
            ws = wb["Первичные"]
            rows = list(ws.iter_rows(values_only=True))
            self.assertEqual(
                rows[0],
                (
                    "Дата","Месяц","Врач","Пациент","Амбулаторная карта","Телефон",
                    "Оплата первичного","Оплатил","Показания ЭВЛК","Показания склеро",
                    "Показания минифлеб","Примечание приёма","Примечание пациента",
                ),
            )
            self.assertEqual(len(rows), 2)
            self.assertEqual(rows[1][1], "Июнь")
            self.assertEqual(rows[1][3], "Пациент Один")
            self.assertEqual(rows[1][6], 4900)
        finally:
            wb.close()

    def test_primary_repeat_export_contains_only_unpaid_primary_who_returned(self):
        response = self.client_for(2).get(
            "/nn/reports/export/primary_repeat.xlsx?clinic_id=1&date_from=2026-06-01&date_to=2026-06-30"
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response.mimetype)
        wb = load_workbook(io.BytesIO(response.data), read_only=True, data_only=True)
        try:
            ws = wb["Первичный + повторные"]
            rows = list(ws.iter_rows(values_only=True))
            self.assertEqual(
                rows[0],
                (
                    "Пациент",
                    "Амбулаторная карта",
                    "Телефон",
                    "Врач первичного",
                    "Дата первичного",
                    "Первый последующий повторный",
                    "Врач повторного",
                    "Всего последующих повторных",
                ),
            )
            self.assertEqual(len(rows), 2)
            self.assertEqual(rows[1][0], "Пациент Два")
            self.assertEqual(rows[1][1], "2")
            self.assertEqual(rows[1][2], "7000000002")
            self.assertEqual(rows[1][3], "Врач Б")
            self.assertEqual(rows[1][4], "2026-06-02")
            self.assertEqual(rows[1][5], "2026-07-02")
            self.assertEqual(rows[1][6], "Врач Б")
            self.assertEqual(rows[1][7], 1)
        finally:
            wb.close()

    def test_treatment_export_contains_patient_detail(self):
        response = self.client_for(2).get(
            "/nn/reports/export/treatment.xlsx?clinic_id=1&year=2026&month=6"
        )
        self.assertEqual(response.status_code, 200)
        wb = load_workbook(io.BytesIO(response.data), read_only=True, data_only=True)
        try:
            ws = wb["Лечение"]
            rows = list(ws.iter_rows(values_only=True))
            self.assertEqual(
                rows[0],
                (
                    "Пациент","Амбулаторная карта","Первичный приём","Врач первичного",
                    "Первое лечение","ЭВЛК","Склеро","Минифлеб","Врач ЭВЛК",
                    "Врач склеро","Врач минифлеб","Оплата лечения найдена",
                    "Сумма строк с лечением","Последующих клинических визитов","Основание",
                ),
            )
            self.assertEqual(rows[1][0], "Пациент Один")
            self.assertEqual(rows[1][2], "2026-06-01")
            self.assertEqual(rows[1][3], "Врач А")
            self.assertEqual(rows[1][11], "Да")
        finally:
            wb.close()


if __name__ == "__main__":
    unittest.main()
