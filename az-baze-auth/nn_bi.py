import io
import math
from collections import defaultdict
from datetime import date, datetime, timedelta

from flask import jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

import nn_normalize
import nn_reports
import nn_upload
from app import db


SERVICE_LABELS = {
    "consultation": "Приём / консультация",
    "evlk": "ЭВЛК",
    "sclero": "Склеротерапия",
    "mini": "Минифлебэктомия",
    "hosiery": "Компрессионный трикотаж",
    "analysis": "Анализы",
    "other": "Прочее",
}

AGE_GROUPS = (
    (0, 24, "до 24"),
    (25, 34, "25–34"),
    (35, 44, "35–44"),
    (45, 54, "45–54"),
    (55, 64, "55–64"),
    (65, 200, "65+"),
)

WEEKDAYS = ("Пн", "Вт", "Ср", "Чт", "Пт", "Сб", "Вс")

RELATION_METRICS = (
    ("revenue", "Выручка"),
    ("visits", "Приёмы"),
    ("patients", "Уникальные пациенты"),
    ("avg_visit", "Средняя сумма на приём"),
    ("zero_share", "Доля приёмов с нулевой суммой"),
    ("treatment_conversion", "Конверсия первичный → лечение"),
)

RELATION_DIMS = (
    ("doctor", "Врач"),
    ("age_group", "Возраст"),
    ("gender", "Пол"),
    ("service", "Категория услуги"),
    ("weekday", "День недели"),
    ("time", "Время суток"),
    ("frequency", "Частота визитов"),
    ("source", "Источник пациента"),
    ("primary_doctor", "Врач первичного"),
    ("treatment_doctor", "Врач лечения"),
)


def _money(value):
    return round(float(value or 0), 2)


def _pct(num, den):
    return round(num / den * 100, 1) if den else None


def _valid_clinic(conn, clinic_id):
    return nn_upload._active_clinic(conn, clinic_id)


def _patient_map(payload):
    return {row["key"]: row for row in payload.get("patients", []) if row.get("key")}


def _service_bucket(visit):
    value = visit.get("service_class")
    if value in SERVICE_LABELS:
        return value
    return nn_normalize._service_class({
        "group": "",
        "service": visit.get("services_text", ""),
        "comment": visit.get("note", ""),
    })


def _age(patient, on_date):
    raw = patient.get("dob")
    if not raw or not on_date:
        return None
    try:
        born = date.fromisoformat(raw)
        current = date.fromisoformat(on_date)
    except ValueError:
        return None
    years = current.year - born.year - ((current.month, current.day) < (born.month, born.day))
    return years if 0 <= years <= 120 else None


def _age_group(patient, on_date):
    years = _age(patient, on_date)
    if years is None:
        return "Не указан"
    for low, high, label in AGE_GROUPS:
        if low <= years <= high:
            return label
    return "Не указан"


def _weekday(value):
    try:
        return WEEKDAYS[date.fromisoformat(value).weekday()]
    except Exception:
        return "Не указан"


def _time_bucket(value):
    if not value:
        return "Не указано"
    try:
        hour = int(value[:2])
    except (TypeError, ValueError):
        return "Не указано"
    if hour < 11:
        return "До 11:00"
    if hour < 14:
        return "11:00–14:00"
    if hour < 17:
        return "14:00–17:00"
    return "После 17:00"


def _frequency_label(count):
    if count <= 1:
        return "1 визит"
    if count == 2:
        return "2 визита"
    return "3+ визита"


def _gender_label(value):
    text = str(value or "").strip().lower()
    if text in {"жен", "ж", "женский", "female", "f"}:
        return "Женский"
    if text in {"муж", "м", "мужской", "male", "m"}:
        return "Мужской"
    return "Не указан"


def _doctor_from_treatment(row):
    for kind in ("evlk", "sclero", "mini"):
        doctors = row.get("doctors", {}).get(kind, [])
        if doctors:
            return doctors[0]
    return ""


def _filters(payload):
    start, end, ctx = nn_reports._bounds(payload)
    doctor = request.args.get("doctor", "").strip()
    service = request.args.get("service", "").strip()
    if service not in SERVICE_LABELS:
        service = ""
    return start, end, ctx, doctor, service


def _visits_for(payload, start, end, doctor="", service=""):
    rows = [
        dict(row) for row in payload.get("visits", [])
        if row.get("date") and start <= row["date"] <= end
    ]
    if doctor:
        rows = [row for row in rows if row.get("doctor") == doctor]
    if service:
        rows = [row for row in rows if _service_bucket(row) == service]
    return rows


def _period_model(payload, start, end, doctor="", service=""):
    visits = _visits_for(payload, start, end, doctor, service)
    patient_keys = {row.get("patient_key") for row in visits if row.get("patient_key")}

    primaries = [
        dict(row) for row in payload.get("primaries", [])
        if row.get("date") and start <= row["date"] <= end
    ]
    repeats = [
        dict(row) for row in payload.get("repeats", [])
        if row.get("date") and start <= row["date"] <= end
    ]
    treatments = [
        dict(row) for row in payload.get("treatments", [])
        if row.get("first_date") and start <= row["first_date"] <= end
    ]

    if doctor:
        primaries = [row for row in primaries if row.get("doctor") == doctor]
        repeats = [row for row in repeats if row.get("doctor") == doctor]
        treatments = [
            row for row in treatments
            if doctor in {
                d for values in row.get("doctors", {}).values() for d in values
            }
        ]

    if doctor or service:
        primaries = [row for row in primaries if row.get("patient_key") in patient_keys]
        repeats = [row for row in repeats if row.get("patient_key") in patient_keys]
        treatments = [row for row in treatments if row.get("patient_key") in patient_keys]

    return {
        "visits": visits,
        "patient_keys": patient_keys,
        "primaries": primaries,
        "repeats": repeats,
        "treatments": treatments,
    }


def _treatment_after_primary(payload, primary):
    pkey = primary.get("patient_key")
    pdate = primary.get("date") or ""
    candidates = [
        row for row in payload.get("treatments", [])
        if row.get("patient_key") == pkey and (row.get("first_date") or "") >= pdate
    ]
    return min(candidates, key=lambda row: row.get("first_date") or "") if candidates else None


def _basic_metrics(payload, model):
    visits = model["visits"]
    pkeys = model["patient_keys"]
    primaries = model["primaries"]
    treatments = model["treatments"]
    revenue = sum(_money(row.get("amount")) for row in visits)
    converted = sum(1 for row in primaries if _treatment_after_primary(payload, row))
    treatment_revenue = sum(
        _money(row.get("amount")) for row in visits
        if _service_bucket(row) in {"evlk", "sclero", "mini"}
    )
    return {
        "patients": len(pkeys),
        "visits": len(visits),
        "revenue": round(revenue, 2),
        "avg_patient": round(revenue / len(pkeys), 2) if pkeys else None,
        "avg_visit": round(revenue / len(visits), 2) if visits else None,
        "primaries": len(primaries),
        "treatments": len({row.get("patient_key") for row in treatments if row.get("patient_key")}),
        "conversion": _pct(converted, len(primaries)),
        "treatment_revenue_share": _pct(treatment_revenue, revenue),
        "zero_visits": sum(1 for row in visits if _money(row.get("amount")) == 0),
        "zero_share": _pct(sum(1 for row in visits if _money(row.get("amount")) == 0), len(visits)),
    }


def _previous_period(payload, start, end, doctor, service):
    try:
        current_start = date.fromisoformat(start)
        current_end = date.fromisoformat(end)
    except Exception:
        return None
    days = (current_end - current_start).days + 1
    prev_end = current_start - timedelta(days=1)
    prev_start = prev_end - timedelta(days=days - 1)
    available = nn_reports._context(payload).get("earliest_date")
    if available and prev_end.isoformat() < available:
        return None
    if available:
        prev_start = max(prev_start, date.fromisoformat(available))
    pstart, pend = prev_start.isoformat(), prev_end.isoformat()
    model = _period_model(payload, pstart, pend, doctor, service)
    return {
        "period": {"start": pstart, "end": pend},
        "metrics": _basic_metrics(payload, model),
    }


def _drivers(current, previous):
    if not previous:
        return []
    prev = previous["metrics"]
    items = []
    for key, label, unit in (
        ("patients", "Количество пациентов", "count"),
        ("avg_patient", "Выручка на пациента", "money"),
        ("conversion", "Конверсия первичный → лечение", "percent"),
        ("treatment_revenue_share", "Доля лечения в выручке", "percent"),
    ):
        cv, pv = current.get(key), prev.get(key)
        if cv is None or pv is None:
            continue
        items.append({
            "key": key,
            "label": label,
            "current": cv,
            "previous": pv,
            "delta": round(cv - pv, 2),
            "unit": unit,
        })
    return items


def _funnel(payload, model):
    primaries = model["primaries"]
    indicated = [
        row for row in primaries
        if any(bool(v) for v in row.get("indications", {}).values())
    ]
    treated = []
    controlled = []
    days_to_treatment = []
    repeats = payload.get("repeats", [])
    for primary in primaries:
        treatment = _treatment_after_primary(payload, primary)
        if not treatment:
            continue
        treated.append(primary)
        try:
            days_to_treatment.append(
                (date.fromisoformat(treatment["first_date"]) - date.fromisoformat(primary["date"])).days
            )
        except (KeyError, TypeError, ValueError):
            pass
        if any(
            row.get("patient_key") == primary.get("patient_key")
            and (row.get("date") or "") >= (treatment.get("first_date") or "")
            for row in repeats
        ):
            controlled.append(primary)
    ordered_days = sorted(days_to_treatment)
    median_days = (
        ordered_days[len(ordered_days) // 2] if ordered_days else None
    )
    return {
        "steps": [
            {"key": "primary", "label": "Первичный", "value": len(primaries)},
            {"key": "indicated", "label": "Есть показания / рекомендация", "value": len(indicated)},
            {"key": "booking", "label": "Запись на лечение", "value": None, "note": "Отдельного надёжного признака в текущих источниках нет"},
            {"key": "treated", "label": "Подтверждено лечение", "value": len(treated)},
            {"key": "control", "label": "Есть последующий контроль / повтор", "value": len(controlled)},
        ],
        "conversion_primary_to_treatment": _pct(len(treated), len(primaries)),
        "conversion_indicated_to_treatment": _pct(
            sum(1 for row in indicated if _treatment_after_primary(payload, row)),
            len(indicated),
        ),
        "days_to_treatment": {
            "average": round(sum(days_to_treatment) / len(days_to_treatment), 1) if days_to_treatment else None,
            "median": median_days,
            "count": len(days_to_treatment),
        },
    }


def _doctor_table(payload, model):
    patients = _patient_map(payload)
    visits = model["visits"]
    primaries = model["primaries"]
    repeats = model["repeats"]
    doctors = sorted({row.get("doctor") for row in visits if row.get("doctor")})
    result = []
    for doctor in doctors:
        dvisits = [row for row in visits if row.get("doctor") == doctor]
        pkeys = {row.get("patient_key") for row in dvisits if row.get("patient_key")}
        revenue = sum(_money(row.get("amount")) for row in dvisits)
        dprim = [row for row in primaries if row.get("doctor") == doctor]
        converted = sum(1 for row in dprim if _treatment_after_primary(payload, row))
        treatment_types = {
            kind: sum(
                1 for row in payload.get("treatments", [])
                if doctor in row.get("doctors", {}).get(kind, [])
                and row.get("patient_key") in pkeys
            )
            for kind in ("evlk", "sclero", "mini")
        }
        result.append({
            "doctor": doctor,
            "patients": len(pkeys),
            "visits": len(dvisits),
            "revenue": round(revenue, 2),
            "avg_patient": round(revenue / len(pkeys), 2) if pkeys else None,
            "avg_visit": round(revenue / len(dvisits), 2) if dvisits else None,
            "primaries": len(dprim),
            "conversion": _pct(converted, len(dprim)),
            "evlk": treatment_types["evlk"],
            "sclero": treatment_types["sclero"],
            "mini": treatment_types["mini"],
            "repeat_visits": sum(1 for row in repeats if row.get("doctor") == doctor),
            "zero_visits": sum(1 for row in dvisits if _money(row.get("amount")) == 0),
        })
    return sorted(result, key=lambda row: (-row["revenue"], row["doctor"]))


def _doctor_transitions(payload, model):
    primary_by_patient = {
        row.get("patient_key"): row
        for row in model["primaries"]
        if row.get("patient_key")
    }
    transitions = defaultdict(int)
    for treatment in model["treatments"]:
        key = treatment.get("patient_key")
        primary = primary_by_patient.get(key)
        if not primary:
            continue
        from_doctor = primary.get("doctor") or "Не указан"
        to_doctor = _doctor_from_treatment(treatment) or "Не указан"
        transitions[(from_doctor, to_doctor)] += 1
    rows = [
        {"primary_doctor": a, "treatment_doctor": b, "patients": count}
        for (a, b), count in transitions.items()
    ]
    return sorted(rows, key=lambda row: (-row["patients"], row["primary_doctor"], row["treatment_doctor"]))


def _count_groups(values, preferred=None):
    counts = defaultdict(int)
    for value in values:
        counts[value or "Не указан"] += 1
    order = preferred or sorted(counts)
    return [{"label": key, "value": counts[key]} for key in order if counts.get(key)]


def _patient_section(payload, model, start, end):
    patients = _patient_map(payload)
    visits = model["visits"]
    pkeys = model["patient_keys"]
    counts = defaultdict(int)
    revenue = defaultdict(float)
    dates = defaultdict(list)
    for row in visits:
        key = row.get("patient_key")
        if not key:
            continue
        counts[key] += 1
        revenue[key] += _money(row.get("amount"))
        if row.get("date"):
            dates[key].append(row["date"])

    age_groups = [_age_group(patients.get(key, {}), end) for key in pkeys]
    gender = [_gender_label(patients.get(key, {}).get("gender")) for key in pkeys]
    frequency = [_frequency_label(counts[key]) for key in pkeys]

    # Cohorts are defined by the first visit in all available history.
    all_visits = defaultdict(list)
    for row in payload.get("visits", []):
        if row.get("patient_key") and row.get("date"):
            all_visits[row["patient_key"]].append(row)
    treatment_by_patient = {
        row.get("patient_key"): row for row in payload.get("treatments", [])
        if row.get("patient_key")
    }
    cohorts = defaultdict(lambda: {
        "patients": 0, "treated": 0, "return_30": 0, "return_60": 0, "return_90": 0, "revenue_90": 0.0
    })
    for key, rows in all_visits.items():
        rows = sorted(rows, key=lambda r: (r.get("date") or "", r.get("start") or ""))
        first = rows[0].get("date")
        if not first or first < start or first > end:
            continue
        cohort = cohorts[first[:7]]
        cohort["patients"] += 1
        if key in treatment_by_patient and (treatment_by_patient[key].get("first_date") or "") >= first:
            cohort["treated"] += 1
        first_date = date.fromisoformat(first)
        for horizon in (30, 60, 90):
            limit = (first_date + timedelta(days=horizon)).isoformat()
            if any(first < (r.get("date") or "") <= limit for r in rows[1:]):
                cohort[f"return_{horizon}"] += 1
        limit90 = (first_date + timedelta(days=90)).isoformat()
        cohort["revenue_90"] += sum(
            _money(r.get("amount")) for r in rows
            if first <= (r.get("date") or "") <= limit90
        )

    cohort_rows = []
    for month in sorted(cohorts):
        item = cohorts[month]
        cohort_rows.append({
            "month": month,
            "patients": item["patients"],
            "treatment_conversion": _pct(item["treated"], item["patients"]),
            "return_30": _pct(item["return_30"], item["patients"]),
            "return_60": _pct(item["return_60"], item["patients"]),
            "return_90": _pct(item["return_90"], item["patients"]),
            "revenue_90": round(item["revenue_90"], 2),
        })

    return {
        "age_groups": _count_groups(age_groups, [x[2] for x in AGE_GROUPS] + ["Не указан"]),
        "gender": _count_groups(gender, ["Женский", "Мужской", "Не указан"]),
        "frequency": _count_groups(frequency, ["1 визит", "2 визита", "3+ визита"]),
        "cohorts": cohort_rows,
        "top_value": [
            {
                "patient_key": key,
                "patient": patients.get(key, {}).get("name", ""),
                "visits": counts[key],
                "revenue": round(revenue[key], 2),
                "last_visit": max(dates[key]) if dates[key] else "",
            }
            for key in sorted(pkeys, key=lambda k: (-revenue[k], patients.get(k, {}).get("name", "")))[:20]
        ],
    }


def _service_section(payload, model):
    visits = model["visits"]
    grouped = defaultdict(lambda: {"visits": 0, "patients": set(), "revenue": 0.0})
    for row in visits:
        bucket = _service_bucket(row)
        item = grouped[bucket]
        item["visits"] += 1
        if row.get("patient_key"):
            item["patients"].add(row["patient_key"])
        item["revenue"] += _money(row.get("amount"))
    categories = []
    for key, item in grouped.items():
        categories.append({
            "key": key,
            "label": SERVICE_LABELS.get(key, key),
            "visits": item["visits"],
            "patients": len(item["patients"]),
            "revenue": round(item["revenue"], 2),
            "avg_visit": round(item["revenue"] / item["visits"], 2) if item["visits"] else None,
        })
    categories.sort(key=lambda row: -row["revenue"])

    control = []
    for row in payload.get("service_control_rows", []):
        control.append({
            "service": row.get("service", ""),
            "qty": row.get("qty", 0),
            "gross": _money(row.get("gross_amount") or row.get("amount")),
            "discount": _money(row.get("discount")),
            "net": _money(row.get("amount")),
        })
    control.sort(key=lambda row: -row["net"])
    chains = defaultdict(lambda: {"patients": 0, "revenue": 0.0})
    by_patient = defaultdict(list)
    for row in visits:
        if row.get("patient_key"):
            by_patient[row["patient_key"]].append(row)
    for rows in by_patient.values():
        rows.sort(key=lambda row: (row.get("date") or "", row.get("start") or ""))
        sequence = []
        total = 0.0
        for row in rows:
            label = SERVICE_LABELS.get(_service_bucket(row), "Прочее")
            if not sequence or sequence[-1] != label:
                sequence.append(label)
            total += _money(row.get("amount"))
        if sequence:
            key = " → ".join(sequence)
            chains[key]["patients"] += 1
            chains[key]["revenue"] += total
    chain_rows = [
        {"chain": key, "patients": item["patients"], "revenue": round(item["revenue"], 2)}
        for key, item in chains.items()
    ]
    chain_rows.sort(key=lambda row: (-row["patients"], -row["revenue"], row["chain"]))

    return {
        "categories": categories,
        "chains": chain_rows[:20],
        "control_top": control[:25],
        "control_note": "Контрольная детализация услуг относится к актуальному загруженному набору и не применяется как отдельный источник оборота.",
    }


def _booking_section(payload, model, start, end, doctor):
    visits = model["visits"]
    by_weekday = defaultdict(lambda: {"visits": 0, "patients": set(), "revenue": 0.0})
    by_time = defaultdict(lambda: {"visits": 0, "patients": set(), "revenue": 0.0})
    lead_days = []
    for row in visits:
        w = _weekday(row.get("date"))
        t = _time_bucket(row.get("start"))
        for bucket, key in ((by_weekday, w), (by_time, t)):
            bucket[key]["visits"] += 1
            if row.get("patient_key"):
                bucket[key]["patients"].add(row["patient_key"])
            bucket[key]["revenue"] += _money(row.get("amount"))
        created = row.get("created_at")
        if created and row.get("date"):
            try:
                created_date = datetime.fromisoformat(created).date()
                appointment = date.fromisoformat(row["date"])
                if appointment >= created_date:
                    lead_days.append((appointment - created_date).days)
            except ValueError:
                pass

    def pack(source, order):
        rows = []
        for key in order:
            item = source.get(key)
            if not item:
                continue
            rows.append({
                "label": key,
                "visits": item["visits"],
                "patients": len(item["patients"]),
                "revenue": round(item["revenue"], 2),
            })
        return rows

    deleted = [
        dict(row) for row in payload.get("deleted_appointments", [])
        if row.get("appointment_date") and start <= row["appointment_date"] <= end
    ]
    if doctor:
        deleted = [row for row in deleted if row.get("doctor") == doctor]
    reasons = _count_groups([row.get("reason") or "Причина недоступна в текущей нормализации" for row in deleted])
    deleted_doctors = _count_groups([row.get("doctor") or "Врач не определён" for row in deleted])

    return {
        "weekday": pack(by_weekday, list(WEEKDAYS)),
        "time": pack(by_time, ["До 11:00", "11:00–14:00", "14:00–17:00", "После 17:00", "Не указано"]),
        "lead_time": {
            "available": bool(lead_days),
            "avg_days": round(sum(lead_days) / len(lead_days), 1) if lead_days else None,
            "median_days": sorted(lead_days)[len(lead_days)//2] if lead_days else None,
            "note": "" if lead_days else "Время создания приёма появится после нормализации источника с новым полем.",
        },
        "deleted_total": len(deleted),
        "deleted_reasons": reasons,
        "deleted_doctors": deleted_doctors,
    }


def _trend_rows(model, start, end):
    visits = model["visits"]
    try:
        span = (date.fromisoformat(end) - date.fromisoformat(start)).days + 1
    except ValueError:
        span = 0
    monthly = span > 62
    grouped = defaultdict(lambda: {"visits": 0, "patients": set(), "revenue": 0.0})
    for row in visits:
        day = row.get("date") or ""
        key = day[:7] if monthly else day
        if not key:
            continue
        item = grouped[key]
        item["visits"] += 1
        if row.get("patient_key"):
            item["patients"].add(row["patient_key"])
        item["revenue"] += _money(row.get("amount"))
    return [
        {
            "period": key,
            "visits": item["visits"],
            "patients": len(item["patients"]),
            "revenue": round(item["revenue"], 2),
        }
        for key, item in sorted(grouped.items())
    ]


def _discount_control(payload):
    rows = payload.get("service_control_rows", [])
    gross = sum(_money(row.get("gross_amount") or row.get("amount")) for row in rows)
    discount = sum(_money(row.get("discount")) for row in rows)
    net = sum(_money(row.get("amount")) for row in rows)
    negative = [
        {
            "service": row.get("service", ""),
            "qty": row.get("qty", 0),
            "amount": _money(row.get("amount")),
        }
        for row in rows
        if float(row.get("qty") or 0) < 0 or _money(row.get("amount")) < 0
    ]
    return {
        "available": bool(rows),
        "gross": round(gross, 2),
        "discount": round(discount, 2),
        "net": round(net, 2),
        "discount_share": _pct(discount, gross),
        "negative_rows": negative,
        "note": "Скидки рассчитаны по актуальному контрольному отчёту услуг; этот источник не используется для повторного суммирования оборота.",
    }


def _marketing_section(payload, model):
    patients = _patient_map(payload)
    visits = model["visits"]
    revenue_by_patient = defaultdict(float)
    for row in visits:
        if row.get("patient_key"):
            revenue_by_patient[row["patient_key"]] += _money(row.get("amount"))
    treatment_keys = {row.get("patient_key") for row in payload.get("treatments", [])}
    grouped = defaultdict(lambda: {"patients": set(), "revenue": 0.0, "treated": set()})
    for key in model["patient_keys"]:
        source = (patients.get(key, {}).get("source") or "").strip()
        if not source:
            continue
        item = grouped[source]
        item["patients"].add(key)
        item["revenue"] += revenue_by_patient[key]
        if key in treatment_keys:
            item["treated"].add(key)
    rows = []
    for source, item in grouped.items():
        count = len(item["patients"])
        rows.append({
            "source": source,
            "patients": count,
            "revenue": round(item["revenue"], 2),
            "avg_patient": round(item["revenue"] / count, 2) if count else None,
            "treatment_conversion": _pct(len(item["treated"]), count),
        })
    rows.sort(key=lambda row: -row["revenue"])
    return {
        "available": bool(rows),
        "sources": rows,
        "note": "" if rows else "Поле «Источник информации о клинике» в текущей нормализованной истории не заполнено. После следующей загрузки новый парсер сохранит его автоматически.",
        "cost_metrics_available": False,
        "cost_note": "CAC/ROAS не рассчитываются без отдельного источника рекламных затрат.",
    }


def _summary_data(conn, clinic_id):
    clinic = _valid_clinic(conn, clinic_id)
    if not clinic:
        raise PermissionError("clinic")
    payload = nn_reports.effective_payload(conn, clinic_id)
    if not payload:
        return {"empty": True, "clinic_id": clinic_id, "clinic_name": clinic["name"]}

    start, end, ctx, doctor, service = _filters(payload)
    model = _period_model(payload, start, end, doctor, service)
    metrics = _basic_metrics(payload, model)
    previous = _previous_period(payload, start, end, doctor, service)
    service_options = [
        {"key": key, "label": label}
        for key, label in SERVICE_LABELS.items()
        if any(_service_bucket(row) == key for row in payload.get("visits", []))
    ]

    return {
        "empty": False,
        "clinic_id": clinic_id,
        "clinic_name": clinic["name"],
        "period": {"start": start, "end": end},
        "filters": {"doctor": doctor, "service": service},
        "context": {
            **ctx,
            "service_options": service_options,
            "relation_metrics": [{"key": k, "label": l} for k, l in RELATION_METRICS],
            "relation_dimensions": [{"key": k, "label": l} for k, l in RELATION_DIMS],
        },
        "overview": {
            "metrics": metrics,
            "previous": previous,
            "drivers": _drivers(metrics, previous),
            "trend": _trend_rows(model, start, end),
        },
        "funnel": _funnel(payload, model),
        "doctors": _doctor_table(payload, model),
        "doctor_transitions": _doctor_transitions(payload, model),
        "patients": _patient_section(payload, model, start, end),
        "services": _service_section(payload, model),
        "bookings": _booking_section(payload, model, start, end, doctor),
        "revenue": {
            "metrics": metrics,
            "discount_control": _discount_control(payload),
            "payment_methods": {
                "available": False,
                "note": "Для способов оплаты нужен дополнительный необязательный источник «Оказанные врачами услуги.pdf». Он не добавлен в пять обязательных файлов.",
            },
        },
        "marketing": _marketing_section(payload, model),
    }


def _relation_event_rows(payload, start, end, doctor, service, metric):
    patients = _patient_map(payload)
    model = _period_model(payload, start, end, doctor, service)
    visits = model["visits"]
    freq = defaultdict(int)
    for row in visits:
        if row.get("patient_key"):
            freq[row["patient_key"]] += 1

    primary_by_patient = {}
    for row in payload.get("primaries", []):
        if row.get("patient_key"):
            primary_by_patient[row["patient_key"]] = row
    treatment_by_patient = {}
    for row in payload.get("treatments", []):
        if row.get("patient_key"):
            treatment_by_patient[row["patient_key"]] = row

    def dims_for(patient_key, doctor_name, day, start_time, service_key):
        patient = patients.get(patient_key, {})
        primary = primary_by_patient.get(patient_key, {})
        treatment = treatment_by_patient.get(patient_key, {})
        return {
            "doctor": doctor_name or "Не указан",
            "age_group": _age_group(patient, end),
            "gender": _gender_label(patient.get("gender")),
            "service": SERVICE_LABELS.get(service_key, service_key or "Не указана"),
            "weekday": _weekday(day),
            "time": _time_bucket(start_time),
            "frequency": _frequency_label(freq.get(patient_key, 0)),
            "source": patient.get("source") or "Не указан",
            "primary_doctor": primary.get("doctor") or "Не указан",
            "treatment_doctor": _doctor_from_treatment(treatment) or "Не указан",
        }

    rows = []
    if metric == "treatment_conversion":
        for primary in model["primaries"]:
            pkey = primary.get("patient_key")
            matching = next(
                (
                    row for row in payload.get("visits", [])
                    if row.get("patient_key") == pkey
                    and row.get("date") == primary.get("date")
                    and (not primary.get("doctor") or row.get("doctor") == primary.get("doctor"))
                ),
                {},
            )
            d = dims_for(
                pkey,
                primary.get("doctor"),
                primary.get("date"),
                matching.get("start"),
                _service_bucket(matching) if matching else "consultation",
            )
            rows.append({
                **d,
                "patient_key": pkey,
                "treated": bool(_treatment_after_primary(payload, primary)),
                "amount": _money(primary.get("amount")),
            })
        return rows

    for visit in visits:
        pkey = visit.get("patient_key")
        rows.append({
            **dims_for(
                pkey,
                visit.get("doctor"),
                visit.get("date"),
                visit.get("start"),
                _service_bucket(visit),
            ),
            "patient_key": pkey,
            "amount": _money(visit.get("amount")),
            "zero": _money(visit.get("amount")) == 0,
        })
    return rows


def _aggregate_relation(rows, metric, dim1, dim2=""):
    if dim1 not in {key for key, _ in RELATION_DIMS}:
        dim1 = "doctor"
    if dim2 not in {key for key, _ in RELATION_DIMS} or dim2 == dim1:
        dim2 = ""

    groups = defaultdict(list)
    for row in rows:
        key = (row.get(dim1) or "Не указан", row.get(dim2) or "") if dim2 else (row.get(dim1) or "Не указан",)
        groups[key].append(row)

    def value(items):
        if metric == "revenue":
            return round(sum(row.get("amount", 0) for row in items), 2)
        if metric == "visits":
            return len(items)
        if metric == "patients":
            return len({row.get("patient_key") for row in items if row.get("patient_key")})
        if metric == "avg_visit":
            return round(sum(row.get("amount", 0) for row in items) / len(items), 2) if items else None
        if metric == "zero_share":
            return _pct(sum(1 for row in items if row.get("zero")), len(items))
        if metric == "treatment_conversion":
            return _pct(sum(1 for row in items if row.get("treated")), len(items))
        return 0

    if not dim2:
        result = [
            {"label": key[0], "value": value(items), "count": len(items)}
            for key, items in groups.items()
        ]
        result.sort(key=lambda row: (-(row["value"] or 0), row["label"]))
        return {"mode": "bar", "rows": result}

    x_values = sorted({key[0] for key in groups})
    y_values = sorted({key[1] for key in groups})
    matrix = []
    for y in y_values:
        matrix.append([
            value(groups.get((x, y), [])) if (x, y) in groups else None
            for x in x_values
        ])
    return {"mode": "heatmap", "x": x_values, "y": y_values, "matrix": matrix}


def relations_data(conn, clinic_id):
    clinic = _valid_clinic(conn, clinic_id)
    if not clinic:
        raise PermissionError("clinic")
    payload = nn_reports.effective_payload(conn, clinic_id)
    if not payload:
        return {"empty": True}
    start, end, _ctx, doctor, service = _filters(payload)
    metric = request.args.get("metric", "revenue")
    if metric not in {key for key, _ in RELATION_METRICS}:
        metric = "revenue"
    dim1 = request.args.get("dim1", "doctor")
    dim2 = request.args.get("dim2", "")
    rows = _relation_event_rows(payload, start, end, doctor, service, metric)
    data = _aggregate_relation(rows, metric, dim1, dim2)
    return {
        "empty": False,
        "period": {"start": start, "end": end},
        "metric": metric,
        "dim1": dim1,
        "dim2": dim2,
        **data,
    }


def bi_page():
    conn = db()
    clinics = nn_upload._clinic_rows(conn)
    return render_template(
        "nn_bi.html",
        clinics=clinics,
        service_labels=SERVICE_LABELS,
    )


def api_summary():
    try:
        clinic_id = int(request.args.get("clinic_id", ""))
    except ValueError:
        return jsonify(error="Укажите клинику."), 400
    try:
        return jsonify(_summary_data(db(), clinic_id))
    except PermissionError:
        return jsonify(error="Клиника недоступна."), 403


def api_relations():
    try:
        clinic_id = int(request.args.get("clinic_id", ""))
    except ValueError:
        return jsonify(error="Укажите клинику."), 400
    try:
        return jsonify(relations_data(db(), clinic_id))
    except PermissionError:
        return jsonify(error="Клиника недоступна."), 403


def _sheet_header(ws, headers):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def _autosize(ws, max_width=34):
    for column in ws.columns:
        width = min(max((len(str(cell.value or "")) for cell in column), default=0) + 2, max_width)
        ws.column_dimensions[column[0].column_letter].width = max(10, width)


def export_xlsx():
    try:
        clinic_id = int(request.args.get("clinic_id", ""))
    except ValueError:
        return "Укажите клинику.", 400
    try:
        data = _summary_data(db(), clinic_id)
    except PermissionError:
        return "Клиника недоступна.", 403
    if data.get("empty"):
        return "Нет данных.", 404

    wb = Workbook()
    ws = wb.active
    ws.title = "Обзор"
    _sheet_header(ws, ["Показатель", "Значение"])
    labels = (
        ("patients", "Уникальные пациенты"),
        ("visits", "Приёмы"),
        ("revenue", "Выручка"),
        ("avg_patient", "Выручка на пациента"),
        ("avg_visit", "Средняя сумма на приём"),
        ("primaries", "Первичные"),
        ("treatments", "Пациенты с лечением"),
        ("conversion", "Конверсия первичный → лечение, %"),
        ("treatment_revenue_share", "Доля лечения в выручке, %"),
        ("zero_share", "Доля приёмов с нулевой суммой, %"),
    )
    for key, label in labels:
        ws.append([label, data["overview"]["metrics"].get(key)])
    ws.append(["Период", f'{data["period"]["start"]} — {data["period"]["end"]}'])
    ws.append(["Клиника", data["clinic_name"]])
    _autosize(ws)

    ws = wb.create_sheet("Врачи")
    headers = ["Врач","Пациенты","Приёмы","Выручка","На пациента","На приём","Первичные","Конверсия %","ЭВЛК","Склеро","Минифлеб","Повторные визиты","Нулевые визиты"]
    _sheet_header(ws, headers)
    for row in data["doctors"]:
        ws.append([row["doctor"],row["patients"],row["visits"],row["revenue"],row["avg_patient"],row["avg_visit"],row["primaries"],row["conversion"],row["evlk"],row["sclero"],row["mini"],row["repeat_visits"],row["zero_visits"]])
    _autosize(ws)

    ws = wb.create_sheet("Когорты")
    _sheet_header(ws, ["Месяц первого визита","Пациенты","Конверсия в лечение %","Возврат 30 дней %","Возврат 60 дней %","Возврат 90 дней %","Выручка 90 дней"])
    for row in data["patients"]["cohorts"]:
        ws.append([row["month"],row["patients"],row["treatment_conversion"],row["return_30"],row["return_60"],row["return_90"],row["revenue_90"]])
    _autosize(ws)

    ws = wb.create_sheet("Услуги")
    _sheet_header(ws, ["Категория","Приёмы","Пациенты","Выручка","Средняя сумма"])
    for row in data["services"]["categories"]:
        ws.append([row["label"],row["visits"],row["patients"],row["revenue"],row["avg_visit"]])
    _autosize(ws)

    ws = wb.create_sheet("Запись и потери")
    _sheet_header(ws, ["Разрез","Значение","Приёмы/количество","Пациенты","Выручка"])
    for row in data["bookings"]["weekday"]:
        ws.append(["День недели",row["label"],row["visits"],row["patients"],row["revenue"]])
    for row in data["bookings"]["time"]:
        ws.append(["Время",row["label"],row["visits"],row["patients"],row["revenue"]])
    for row in data["bookings"]["deleted_reasons"]:
        ws.append(["Причина удаления",row["label"],row["value"],None,None])
    _autosize(ws)

    ws = wb.create_sheet("Маркетинг")
    _sheet_header(ws, ["Источник","Пациенты","Выручка","Выручка на пациента","Конверсия в лечение %"])
    for row in data["marketing"]["sources"]:
        ws.append([row["source"],row["patients"],row["revenue"],row["avg_patient"],row["treatment_conversion"]])
    if not data["marketing"]["sources"]:
        ws.append(["Источник не заполнен в текущей нормализованной истории",None,None,None,None])
    _autosize(ws)

    output = io.BytesIO()
    wb.save(output)
    wb.close()
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name=f"varikoza-net-kz-bi-{clinic_id}-{data['period']['start']}-{data['period']['end']}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
